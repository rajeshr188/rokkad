import io
from datetime import date

from django.contrib.auth.decorators import login_required
from django.core.exceptions import ValidationError, ObjectDoesNotExist
from django.core.paginator import Paginator
from django.db import IntegrityError
from django.http import HttpResponse
from django.shortcuts import render, redirect
from django.views.decorators.cache import never_cache
from django.views.decorators.http import require_http_methods, require_GET

from apps.tenant_apps.loans.services.history_setup import require_history_setup_access
from . import guided_openings as service, opening_register as register
from .guided_opening_forms import UploadRegisterForm, ReviewRegisterForm
from .models import GuidedOpeningBatch
from .parsers import MAX_BYTES


@login_required
@never_cache
@require_http_methods(['GET', 'POST'])
def upload(request):
    args = {'workspace_id': request.workspace.pk, 'actor': request.user}
    require_history_setup_access(**args, read_only=request.method == 'GET')
    form = UploadRegisterForm(request.POST if request.method == 'POST' else None, request.FILES or None)
    if request.method == 'POST' and form.is_valid():
        file = form.cleaned_data['file']
        try:
            batch = service.stage(**args, source_key=form.cleaned_data['source_key'], content=file.read(MAX_BYTES + 1), filename=file.name)
            return redirect('workspace_portability:guided_batch', workspace_slug=request.workspace.slug, batch_id=batch.public_id)
        except ValueError as exc:
            form.add_error(None, str(exc))
    batches = Paginator(GuidedOpeningBatch.objects.filter(workspace=request.workspace).order_by('-created_at', '-pk'), 20).get_page(request.GET.get('page'))
    return render(request, 'data_portability/guided_upload.html', {'form': form, 'batches': batches})


@login_required
@never_cache
@require_http_methods(['GET', 'POST'])
def review(request, batch_id):
    args = {'workspace_id': request.workspace.pk, 'actor': request.user, 'batch_id': batch_id}
    batch = service.get_batch(**args)
    form, approval, error = None, None, None
    if batch.state in {'STAGED', 'READY'}:
        form = ReviewRegisterForm(request.POST if request.method == 'POST' and request.POST.get('action') == 'preview' else None,
            workspace=request.workspace, batch=batch)
    if request.method == 'POST':
        try:
            action = request.POST.get('action')
            if action == 'preview' and form is not None and form.is_valid():
                batch, approval = service.preview(**args, mapping=form.mapping())
            elif action == 'commit':
                service.commit(**args, approval=request.POST.get('approval', ''), confirmed=request.POST.get('confirmed') == 'yes')
                return redirect('workspace_portability:guided_batch', workspace_slug=request.workspace.slug, batch_id=batch.public_id)
            elif action == 'cancel':
                service.cancel(**args, confirmed=request.POST.get('confirmed') == 'yes')
                return redirect('workspace_portability:guided_batch', workspace_slug=request.workspace.slug, batch_id=batch.public_id)
            elif action != 'preview':
                error = 'Choose a supported import action.'
        except (ValueError, ValidationError, ObjectDoesNotExist) as exc:
            error = str(exc)
        except IntegrityError:
            error = 'The destination changed or contains a conflicting record. Nothing was imported; refresh and review again.'
    return render(request, 'data_portability/guided_review.html', {'batch': batch, 'form': form,
        'approval': approval, 'error': error, 'unfinished': batch.state in {'STAGED', 'READY'},
        'cutover_date': date.fromisoformat(batch.mapping['settings']['cutover']) if batch.mapping else None})


def template_bytes():
    # The application already uses openpyxl for its hardened values-only reader.
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Outstanding loans'
    sheet.append(register.COLUMNS)
    sheet.freeze_panes = 'D2'
    sheet.row_dimensions[1].height = 42
    for col in range(1, len(register.COLUMNS) + 1):
        sheet.column_dimensions[get_column_letter(col)].width = 22 if col in (4, 9) else 18
        cell = sheet.cell(1, col)
        cell.fill, cell.font = PatternFill('solid', fgColor='17365D'), Font(color='FFFFFF', bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical='center')
        for row in range(2, 22):
            # Text format preserves identifiers/date text and matches the strict parser.
            sheet.cell(row, col).number_format = '@'
            sheet.cell(row, col).alignment = Alignment(vertical='top')
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


@login_required
@never_cache
@require_GET
def template(request):
    require_history_setup_access(request.workspace.pk, request.user, read_only=True)
    response = HttpResponse(template_bytes(), content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename="rokkad-outstanding-loans.xlsx"'
    return response


@login_required
@never_cache
@require_GET
def guide(request):
    require_history_setup_access(request.workspace.pk, request.user, read_only=True)
    return render(request, 'data_portability/guided_guide.html')
