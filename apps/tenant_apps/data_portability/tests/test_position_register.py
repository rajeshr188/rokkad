"""Fictional essential-position registers, not customer-book attestations."""
import copy
import csv
import io
from datetime import date
from types import SimpleNamespace
from unittest.mock import patch
import importlib

from django.core.exceptions import PermissionDenied
from django.core.files.uploadedfile import SimpleUploadedFile
from django.db import DatabaseError, transaction, connection
from django.test import Client, override_settings
from django.urls import reverse
from django.utils import timezone

from apps.tenancy.testing import WorkspaceTestCase
from apps.tenant_apps.loans import models as m
from apps.tenant_apps.loans.services.closed_position import read_position
from apps.tenant_apps.loans.services.pawn_repayment import record_pawn_loan_repayment
from apps.tenant_apps.loans.services.pawn_reversal import reverse_pawn_loan_event
from apps.tenant_apps.loans.tests.test_opening_checkpoint import balance_values
from apps.tenant_apps.party.models import Party
from apps.tenant_apps.data_portability import guided_openings as service, opening_register as register
from apps.tenant_apps.data_portability.guided_opening_views import template_bytes
from apps.tenant_apps.data_portability.models import GuidedOpeningBatch
from apps.tenant_apps.data_portability.parsers import PortabilityError, parse_source
from .test_guided_openings import GuidedOpeningTests, row


def checkpoint_row(**changes):
    values = row(outstanding_principal='800', unpaid_interest='5')
    del values['first_month_paid']
    values.update(period_principal_base='900', recognized_interest='9',
        current_period_recognized_interest='9', current_period_unpaid_interest='5',
        current_period_advance_interest='0', future_advance_interest='0', interest_quantum='0.01')
    values.update(changes)
    return values


def closed_row(**changes):
    values = dict(loan_ref='old-1', loan_number='OLD00001', borrower_ref='borrower-1',
        borrower_name='Fictional Borrower', borrower_phone='', original_date='', closed_on='',
        original_principal='', monthly_rate_percent='', tenure_months='',
        outstanding_principal='0', unpaid_interest='0', unpaid_fees='0', collateral_returned='yes')
    values.update(changes)
    return values


def source_csv(rows, profile):
    stream = io.StringIO()
    writer = csv.DictWriter(stream, fieldnames=register.TEMPLATES[profile])
    writer.writeheader()
    writer.writerows(rows)
    return stream.getvalue().encode()


class PositionRegisterTests(GuidedOpeningTests):
    def position_batch(self, rows=None, profile=register.CHECKPOINT_PROFILE):
        return self.staged(content=source_csv(rows or [checkpoint_row()], profile))

    def closed_mapping(self):
        values = copy.deepcopy(self.mapping)
        for key in ('revision_id', 'grace_days'):
            del values['settings'][key]
        return values

    def test_reduced_checkpoint_continues_repayment_reversal_and_original_anniversary(self):
        with self.scoped():
            before = (Party.objects.count(), m.LoanProduct.objects.count())
            batch, token = self.reviewed(self.position_batch())
            self.assertEqual((Party.objects.count(), m.LoanProduct.objects.count()), before)
            self.assertEqual(batch.preview['rows'][0]['checkpoint']['bases'][0]['principal_base'], '900')
            result = self.admit(batch, token)
            loan = m.PawnLoan.objects.get(pk=result.results[0]['loan_id'])
            self.assertEqual(loan.historical_import.document['review']['profile'], 'loan-opening-review/4')
            self.assertEqual(loan.policy_snapshot.policy_version, 2)
            self.assertEqual(balance_values(loan, date(2021, 2, 2)), dict(principal='800', interest='5', fees='0'))
            with patch('django.utils.timezone.localdate', return_value=date(2021, 2, 20)):
                payment = record_pawn_loan_repayment(loan.pk, amount=205, actor=self.actor, request_key='checkpoint-receipt')
            self.assertEqual(balance_values(loan, date(2021, 3, 1))['interest'], '0')
            self.assertEqual(balance_values(loan, date(2021, 3, 2)), dict(principal='600', interest='6', fees='0'))
            with patch('django.utils.timezone.localdate', return_value=date(2021, 3, 3)):
                reverse_pawn_loan_event(payment.loan_event.pk, actor=self.actor, reason='Fictional correction')
            self.assertEqual(balance_values(loan, date(2021, 3, 3)), dict(principal='800', interest='13', fees='0'))
            self.assertEqual(self.admit(batch, token).results, result.results)

    def test_first_month_advance_and_reduction_do_not_start_next_charge_on_anniversary(self):
        with self.scoped():
            values = checkpoint_row(original_principal='10000', outstanding_principal='9000', period_principal_base='10000',
                monthly_rate_percent='2', unpaid_interest='0', recognized_interest='0',
                current_period_recognized_interest='0', current_period_unpaid_interest='0', current_period_advance_interest='200')
            mapping = copy.deepcopy(self.mapping)
            mapping['settings']['cutover'] = '2021-01-20'
            batch, token = self.reviewed(self.position_batch([values]), mapping)
            result = self.admit(batch, token)
            loan = m.PawnLoan.objects.get(pk=result.results[0]['loan_id'])
            self.assertEqual(balance_values(loan, date(2021, 2, 1))['interest'], '0')
            self.assertEqual(balance_values(loan, date(2021, 2, 2))['interest'], '180')
            self.assertEqual(list(loan.loan_events.values_list('event_kind', flat=True)), ['MIGRATION_OPENING'])
            self.assertFalse(m.PawnLoanDisbursalSnapshot.objects.exists())

    def test_checkpoint_missing_or_contradictory_facts_hold_entire_batch(self):
        with self.scoped():
            for changes in ({'period_principal_base': ''}, {'period_principal_base': '700'},
                {'current_period_recognized_interest': '8'}, {'current_period_unpaid_interest': '6'},
                {'recognized_interest': ''}, {'future_advance_interest': '8'}, {'future_advance_interest': ''},
                {'interest_quantum': '0.1'}, {'outstanding_principal': '1100'}):
                with self.subTest(changes=changes):
                    batch = self.position_batch([checkpoint_row(), checkpoint_row(loan_ref='old-2', loan_number='OLD00002', **changes)])
                    batch, token = service.preview(**self.base, batch_id=batch.public_id, mapping=self.mapping)
                    self.assertIsNone(token)
                    self.assertTrue(batch.preview['issues'])
                    self.assertFalse(m.PawnLoan.objects.exists())

    def test_item_bases_and_half_up_rounding_are_reviewed_per_item(self):
        with self.scoped():
            values = dict(original_principal='100', outstanding_principal='90', period_principal_base='100',
                monthly_rate_percent='1.5', recognized_interest='4', current_period_recognized_interest='4',
                unpaid_interest='4', current_period_unpaid_interest='4', interest_quantum='1')
            batch, token = self.reviewed(self.position_batch([checkpoint_row(**values), checkpoint_row(item_ref='item-2', **values)]))
            self.assertEqual(batch.preview['rows'][0]['checkpoint']['expected_period_interest'], '4')
            self.admit(batch, token)
            self.assertEqual(m.PawnCollateralItem.objects.count(), 2)

    def test_closed_register_unknown_terms_nonposting_review_and_exact_replay(self):
        with self.scoped():
            products = m.LoanProduct.objects.count()
            batch = self.position_batch([closed_row()], register.CLOSED_PROFILE)
            batch, token = self.reviewed(batch, self.closed_mapping())
            self.assertFalse(m.PawnLoan.objects.exists())
            self.assertEqual(m.LoanProduct.objects.count(), products)
            result = self.admit(batch, token)
            loan = m.PawnLoan.objects.get(pk=result.results[0]['loan_id'])
            self.assertEqual(loan.state, 'CLOSED')
            self.assertTrue(loan.is_imported_closed_position)
            self.assertIsNone(loan.principal_amount)
            self.assertIsNone(loan.tenure_months)
            self.assertIsNone(loan.policy_snapshot_id)
            self.assertIsNone(loan.product_version_id)
            self.assertEqual(read_position(loan)['position']['custody'], 'RETURNED_TO_BORROWER')
            self.assertEqual(read_position(loan)['earlier_history'], 'UNAVAILABLE')
            self.assertEqual(balance_values(loan, date(2021, 2, 2)), dict(principal='0', interest='0', fees='0'))
            self.assertEqual(list(loan.loan_events.values_list('event_kind', flat=True)), ['MIGRATION_OPENING'])
            self.assertFalse(m.PawnLoanRelease.objects.exists())
            self.assertFalse(m.PawnLoanDisbursalSnapshot.objects.exists())
            self.assertEqual(self.admit(batch, token).results, result.results)
            self.assertEqual(self.position_batch([closed_row()], register.CLOSED_PROFILE).pk, batch.pk)

    def test_closed_known_terms_and_series_number_claim_are_preserved(self):
        with self.scoped():
            values = closed_row(loan_number='H-00003', original_date='01/01/2021', closed_on='01/02/2021',
                original_principal='2000', monthly_rate_percent='2', tenure_months='12')
            batch, token = self.reviewed(self.position_batch([values], register.CLOSED_PROFILE), self.closed_mapping())
            self.sequence.refresh_from_db()
            self.assertEqual(self.sequence.next_number, 1)
            result = self.admit(batch, token)
            loan = m.PawnLoan.objects.get(pk=result.results[0]['loan_id'])
            self.assertEqual((loan.principal_amount, loan.monthly_interest_rate, loan.tenure_months), (2000, 2, 12))
            self.sequence.refresh_from_db()
            self.assertEqual(self.sequence.next_number, 4)

    def test_closed_nonzero_unknown_balances_and_unreturned_custody_are_rejected(self):
        with self.scoped():
            for changes in ({'unpaid_interest': ''}, {'unpaid_fees': '1'}, {'outstanding_principal': '1'},
                {'collateral_returned': ''}, {'collateral_returned': 'no'}, {'original_principal': '0'},
                {'closed_on': '03/02/2021'}, {'original_date': 'invalid'}):
                with self.subTest(changes=changes):
                    batch = self.position_batch([closed_row(**changes)], register.CLOSED_PROFILE)
                    batch, token = service.preview(**self.base, batch_id=batch.public_id, mapping=self.closed_mapping())
                    self.assertIsNone(token)
                    self.assertTrue(batch.preview['issues'])
                    self.assertFalse(m.PawnLoan.objects.exists())

    def test_duplicate_closed_rows_and_changed_source_cannot_overwrite_active_position(self):
        with self.scoped():
            batch = self.position_batch([closed_row(), closed_row()], register.CLOSED_PROFILE)
            batch, token = service.preview(**self.base, batch_id=batch.public_id, mapping=self.closed_mapping())
            self.assertIsNone(token)
            active, token = self.reviewed(self.position_batch())
            self.admit(active, token)
            mapping = self.closed_mapping()
            mapping['borrowers']['borrower-1'] = Party.objects.get(display_name='Fictional Borrower').party_code
            batch = self.position_batch([closed_row()], register.CLOSED_PROFILE)
            batch, token = service.preview(**self.base, batch_id=batch.public_id, mapping=mapping)
            self.assertIsNone(token)
            self.assertEqual(m.PawnLoan.objects.count(), 1)
            self.assertEqual(m.PawnLoan.objects.get().state, 'ACTIVE')

    def test_closed_inactive_customer_match_and_invalid_setup_are_distinct(self):
        with self.scoped():
            party = self.source_party.identity.party
            Party.objects.filter(pk=party.pk).update(status='INACTIVE')
            mapping = self.closed_mapping()
            mapping['borrowers']['borrower-1'] = party.party_code
            batch, token = self.reviewed(self.position_batch([closed_row()], register.CLOSED_PROFILE), mapping)
            self.admit(batch, token)
            self.assertEqual(m.PawnLoan.objects.get().borrower_id, party.pk)
            bad = self.closed_mapping()
            bad['settings']['series_id'] = self.b.pk + 999999
            with self.assertRaises(PortabilityError):
                service.prepare(**self.base, batch_id=self.position_batch([closed_row(loan_ref='new', loan_number='OLD00002')], register.CLOSED_PROFILE).public_id, mapping=bad)

    def test_closed_partial_failure_rolls_back_loans_and_new_customers(self):
        with self.scoped():
            before = Party.objects.count()
            batch = self.position_batch([closed_row(), closed_row(loan_ref='old-2', loan_number='OLD00002')], register.CLOSED_PROFILE)
            batch, token = self.reviewed(batch, self.closed_mapping())
            original = service.admit_closed_position
            def fail_second(**kwargs):
                if kwargs['document']['source']['loan_id'] == 'old-2':
                    raise ValueError('Fictional commit failure')
                return original(**kwargs)
            with patch.object(service, 'admit_closed_position', side_effect=fail_second), self.assertRaises(PortabilityError):
                self.admit(batch, token)
            self.assertFalse(m.PawnLoan.objects.exists())
            self.assertEqual(Party.objects.count(), before)

    def test_database_retains_new_profiles_and_refuses_mismatched_completion(self):
        with self.scoped():
            batch, token = self.reviewed(self.position_batch([closed_row()], register.CLOSED_PROFILE), self.closed_mapping())
            with self.assertRaises(DatabaseError), transaction.atomic():
                GuidedOpeningBatch.objects.filter(pk=batch.pk).update(document={'profile': register.PROFILE, 'rows': batch.document['rows']})
            migration = importlib.import_module('apps.tenant_apps.data_portability.migrations.0021_position_register_profiles')
            with self.assertRaisesMessage(DatabaseError, 'Cannot retire'), transaction.atomic():
                with connection.cursor() as cursor:
                    cursor.execute(migration.REVERSE)
            result = self.admit(batch, token)
            with self.assertRaises(DatabaseError), transaction.atomic():
                GuidedOpeningBatch.objects.filter(pk=batch.pk).update(results=[])
        with self.scoped(self.b):
            self.assertFalse(GuidedOpeningBatch.objects.filter(pk=batch.pk).exists())
            with self.assertRaises(PermissionDenied):
                service.get_batch(workspace_id=self.b.pk, actor=self.actor, batch_id=batch.public_id)

    def test_closed_batch_cannot_claim_an_active_opening_as_its_result(self):
        with self.scoped():
            active, token = self.reviewed(self.position_batch())
            active = self.admit(active, token)
            mapping = self.closed_mapping()
            mapping['borrowers']['borrower-1'] = Party.objects.get(display_name='Fictional Borrower').party_code
            closed, _ = self.reviewed(self.position_batch([closed_row(loan_ref='old-2', loan_number='OLD00002')], register.CLOSED_PROFILE), mapping)
            with self.assertRaisesMessage(DatabaseError, 'Guided results'), transaction.atomic():
                GuidedOpeningBatch.objects.filter(pk=closed.pk).update(state='COMPLETED', results=active.results,
                    committed_by=self.actor, committed_at=timezone.now())

    def test_checkpoint_batch_cannot_claim_legacy_unchanged_principal_review(self):
        with self.scoped():
            legacy, token = self.reviewed()
            legacy = self.admit(legacy, token)
            mapping = copy.deepcopy(self.mapping)
            mapping['borrowers']['borrower-1'] = Party.objects.get(display_name='Fictional Borrower').party_code
            checkpoint, _ = self.reviewed(self.position_batch([checkpoint_row(loan_ref='old-2', loan_number='OLD00002')]), mapping)
            with self.assertRaisesMessage(DatabaseError, 'Guided results'), transaction.atomic():
                GuidedOpeningBatch.objects.filter(pk=checkpoint.pk).update(state='COMPLETED', results=legacy.results,
                    committed_by=self.actor, committed_at=timezone.now())

    def test_zero_remaining_item_preserves_custody_and_continues_only_live_debt(self):
        with self.scoped():
            values = dict(unpaid_interest='18', current_period_unpaid_interest='18',
                recognized_interest='18', current_period_recognized_interest='18')
            batch = self.position_batch([checkpoint_row(outstanding_principal='0', **values), checkpoint_row(item_ref='item-2', **values)])
            batch, token = service.preview(**self.base, batch_id=batch.public_id, mapping=self.mapping)
            self.assertIsNotNone(token, batch.preview['issues'])
            self.assertFalse(m.PawnLoan.objects.exists())
            result = self.admit(batch, token)
            loan = m.PawnLoan.objects.get(pk=result.results[0]['loan_id'])
            self.assertEqual(list(loan.collateral_items.order_by('pk').values_list('allocated_principal','custody_state')),
                [(0,'IN_VAULT'),(800,'IN_VAULT')])
            from apps.tenant_apps.loans.services.pawn_tranches import get_pawn_principal_tranche_balances
            self.assertEqual([row.principal_outstanding for row in get_pawn_principal_tranche_balances(loan)], [0,800])
            self.assertEqual(balance_values(loan,date(2021,2,2)),dict(principal='800',interest='18',fees='0'))
            self.assertEqual(balance_values(loan,date(2021,3,2)),dict(principal='800',interest='26',fees='0'))

    def test_unreviewed_zero_allocation_is_refused_by_database(self):
        with self.scoped():
            batch, token = self.reviewed(self.position_batch())
            result = self.admit(batch,token)
            item = m.PawnCollateralItem.objects.get(loan_id=result.results[0]['loan_id'])
            with self.assertRaisesMessage(DatabaseError,'accepted explicit opening checkpoint'), transaction.atomic():
                m.PawnCollateralItem.objects.filter(pk=item.pk).update(allocated_principal=0)
                with connection.cursor() as cursor:
                    cursor.execute('SET CONSTRAINTS ALL IMMEDIATE')

    def test_zero_retained_item_roundtrip_preserves_both_items_and_future_interest(self):
        from apps.tenant_apps.loans.services.servicing_bundle import export_servicing_bundle,restore_servicing_bundle
        from apps.tenant_apps.loans.services.servicing_bundle_contract import read_bundle
        from apps.tenant_apps.loans.tests.test_servicing_bundle import destination
        from apps.tenant_apps.loans.tests.test_recorded_auctions import on
        with on(date(2021,2,3)):
            with self.scoped():
                values=dict(unpaid_interest='18',current_period_unpaid_interest='18',
                    recognized_interest='18',current_period_recognized_interest='18')
                batch=self.position_batch([checkpoint_row(outstanding_principal='0',**values),checkpoint_row(item_ref='item-2',**values)])
                batch,token=service.preview(**self.base,batch_id=batch.public_id,mapping=self.mapping)
                admitted=self.admit(batch,token)
                loan_id=admitted.results[0]['loan_id']
                content=export_servicing_bundle(workspace_id=self.a.pk,actor=self.actor,loan_id=loan_id)
            document,_=read_bundle(content)
            mapping=destination(self,document)
            with self.scoped(self.b):
                args=dict(workspace_id=self.b.pk,actor=self.actor,content=content,mapping=mapping)
                review=restore_servicing_bundle(**args)
                result=restore_servicing_bundle(**args,expected_sha256=review['sha256'],confirmed=True)
                loan=m.PawnLoan.objects.get(pk=result['loans'][0])
                self.assertEqual(list(loan.collateral_items.order_by('pk').values_list('allocated_principal','custody_state')),[(0,'IN_VAULT'),(800,'IN_VAULT')])
                self.assertEqual(balance_values(loan,date(2021,3,2)),dict(principal='800',interest='26',fees='0'))
                again,_=read_bundle(export_servicing_bundle(workspace_id=self.b.pk,actor=self.actor,loan_id=loan.pk))
                self.assertEqual(again['positions'][str(loan.pk)]['recorded'],document['positions'][str(loan_id)]['recorded'])

    def test_both_templates_roundtrip_and_formulas_are_rejected(self):
        from openpyxl import load_workbook
        for profile, values in ((register.CHECKPOINT_PROFILE, checkpoint_row()), (register.CLOSED_PROFILE, closed_row())):
            workbook = load_workbook(io.BytesIO(template_bytes(profile)))
            sheet = workbook.active
            self.assertEqual(tuple(cell.value for cell in sheet[1]), register.TEMPLATES[profile])
            for col, key in enumerate(register.TEMPLATES[profile], 1):
                sheet.cell(2, col, values[key])
            output = io.BytesIO()
            workbook.save(output)
            with self.scoped():
                batch = self.staged(content=output.getvalue(), filename='positions.xlsx')
                self.assertEqual(batch.document['profile'], profile)
                self.reviewed(batch, self.closed_mapping() if profile == register.CLOSED_PROFILE else self.mapping)
            sheet.cell(2, 1, '=1+1')
            output = io.BytesIO()
            workbook.save(output)
            with self.assertRaises(PortabilityError):
                parse_source(output.getvalue(), 'positions.xlsx')

    @override_settings(STORAGES={'default': {'BACKEND': 'django.core.files.storage.FileSystemStorage'},
        'staticfiles': {'BACKEND': 'django.contrib.staticfiles.storage.StaticFilesStorage'}})
    def test_browser_closed_upload_prepare_review_and_confirm_needs_no_policy_fields(self):
        WorkspaceTestCase.start_active_trial(SimpleNamespace(tenant=self.a))
        client = Client()
        client.force_login(self.actor)
        kwargs = dict(workspace_slug=self.a.slug)
        upload = reverse('workspace_portability:guided_upload', kwargs=kwargs)
        response = client.post(upload, dict(source_key='fictional-ledger',
            file=SimpleUploadedFile('closed.csv', source_csv([closed_row()], register.CLOSED_PROFILE), 'text/csv')))
        self.assertEqual(response.status_code, 302)
        detail = response.url
        response = client.get(detail)
        self.assertNotIn('revision_id', response.context['form'].fields)
        self.assertNotIn('grace_days', response.context['form'].fields)
        data = dict(self.closed_mapping()['settings'], mode_0='NEW', action='prepare')
        response = client.post(detail, data)
        self.assertContains(response, 'No financial admission has been approved')
        data['action'] = 'preview'
        response = client.post(detail, data)
        self.assertContains(response, 'Closed with zero debt and collateral returned')
        approval = response.context['approval']
        self.assertTrue(approval)
        response = client.post(detail, dict(action='commit', approval=approval, confirmed='yes'))
        self.assertEqual(response.status_code, 302)
        self.assertContains(client.get(detail), '1 existing loans imported.')
        template = reverse('workspace_portability:guided_template', kwargs=kwargs)
        for profile in (register.CHECKPOINT_PROFILE, register.CLOSED_PROFILE):
            self.assertEqual(client.get(template, dict(profile=profile)).status_code, 200)
        self.assertEqual(client.get(template, dict(profile='unknown')).status_code, 404)
