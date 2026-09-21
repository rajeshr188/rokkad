import copy
import io
import stat
import zipfile
from xml.etree import ElementTree as ET
from datetime import date
from unittest.mock import patch

from django.core.exceptions import PermissionDenied
from django.test import SimpleTestCase, RequestFactory, override_settings
from django.core.files.uploadedfile import SimpleUploadedFile
from openpyxl import Workbook

from apps.tenant_apps.data_portability import child_contracts as contract, presets, services, views, xlsx
from apps.tenant_apps.data_portability.models import ImportBatch
from apps.tenant_apps.data_portability.parsers import PortabilityError, parse_source
from apps.tenant_apps.party.models import Party, PartyRoleType
from .fixtures import PortabilityFixture
from .test_portability import CSV, MAPPING


def workbook_bytes(rows=None, change=None):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Party register"
    for row in rows or [["Legacy", "Party", "Phone"], ["old-1", "Asha Devi", ""]]:
        sheet.append(row)
    if change:
        change(workbook, sheet)
    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def rewrite(content, changes=None, additions=None):
    output = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(content)) as original, zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as result:
        for part in original.infolist():
            data = original.read(part)
            transform = (changes or {}).get(part.filename)
            result.writestr(part.filename, transform(data) if transform else data)
        for name, data in (additions or {}).items():
            result.writestr(name, data)
    return output.getvalue()


class XlsxParserTests(SimpleTestCase):
    def parse(self, content):
        return parse_source(content, "register.xlsx")

    def test_literal_text_numeric_and_boolean_cells_preserve_values(self):
        content = workbook_bytes([["ID", "Name", "Number", "Flag", "Blank"],
            ["000123", "आशा", 0.125, True, None], ["=literal text", "Next", 42, False, ""]],
            lambda w, s: setattr(s["A3"], "data_type", "s"))
        _, kind, headers, rows = self.parse(content)
        self.assertEqual(kind, "xlsx")
        self.assertEqual(rows[0], (2, dict(zip(headers, ["000123", "आशा", "0.125", "true", ""]))))
        self.assertEqual(rows[1][1]["ID"], "=literal text")

    def test_blank_rows_keep_physical_row_numbers(self):
        content = workbook_bytes([["ID", "Name"], ["a", "First"], [], ["b", "Second"]])
        self.assertEqual([n for n, _ in self.parse(content)[3]], [2, 4])

    def test_excel_shared_strings_preserve_text_and_leading_zeros(self):
        ns = xlsx.SHEET_NS
        strings = ET.Element(ns + "sst")
        def shared(data):
            root = ET.fromstring(data)
            for cell in root.iter(ns + "c"):
                inline = cell.find(ns + "is")
                if inline is None:
                    continue
                value = "".join(n.text or "" for n in inline.iter(ns + "t"))
                cell.remove(inline)
                cell.set("t", "s")
                ET.SubElement(cell, ns + "v").text = str(len(strings))
                ET.SubElement(ET.SubElement(strings, ns + "si"), ns + "t").text = value
            return ET.tostring(root)
        # Build shared strings before adding the part to the archive.
        content = workbook_bytes([["ID", "Name"], ["000123", "आशा"]])
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            sheet = shared(archive.read("xl/worksheets/sheet1.xml"))
        content = rewrite(content, {
            "xl/worksheets/sheet1.xml": lambda d: sheet,
            "xl/_rels/workbook.xml.rels": lambda d: d.replace(b"</Relationships>",
                b'<Relationship Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/sharedStrings" Target="sharedStrings.xml" Id="rIdStrings"/></Relationships>'),
            "[Content_Types].xml": lambda d: d.replace(b"</Types>",
                b'<Override PartName="/xl/sharedStrings.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml"/></Types>')},
            {"xl/sharedStrings.xml": ET.tostring(strings)})
        self.assertEqual(self.parse(content)[3], [(2, {"ID": "000123", "Name": "आशा"})])

    def test_formulas_rejected_even_with_cached_values_before_openpyxl(self):
        content = workbook_bytes(change=lambda w, s: setattr(s["B2"], "value", "=1+1"))
        content = rewrite(content, {"xl/worksheets/sheet1.xml": lambda d: d.replace(b"<v></v>", b"<v>2</v>")})
        with patch.object(xlsx, "load_workbook") as reader, self.assertRaises(PortabilityError):
            self.parse(content)
        reader.assert_not_called()

    def test_hidden_rows_columns_and_multiple_sheets_rejected(self):
        for change in (lambda w, s: setattr(s.row_dimensions[2], "hidden", True),
                       lambda w, s: setattr(s.column_dimensions["B"], "hidden", True),
                       lambda w, s: w.create_sheet("Other")):
            with self.subTest(change=change), self.assertRaises(PortabilityError):
                self.parse(workbook_bytes(change=change))
        content = rewrite(workbook_bytes(), {"xl/workbook.xml": lambda d: d.replace(b'state="visible"', b'state="hidden"')})
        with self.assertRaises(PortabilityError):
            self.parse(content)
        content = rewrite(workbook_bytes(), {"xl/worksheets/sheet1.xml": lambda d: d.replace(b'<row r="2"', b'<row r="2" ht="0"')})
        with self.assertRaises(PortabilityError):
            self.parse(content)

    def test_merges_filters_links_and_errors_rejected(self):
        for change in (lambda w, s: s.merge_cells("B2:C2"),
                       lambda w, s: setattr(s.auto_filter, "ref", "A1:C2"),
                       lambda w, s: setattr(s["B2"], "hyperlink", "https://example.com"),
                       lambda w, s: setattr(s["B2"], "value", "#DIV/0!")):
            with self.subTest(change=change), self.assertRaises(PortabilityError):
                self.parse(workbook_bytes(change=change))

    def test_dates_formatted_numbers_and_excess_precision_rejected(self):
        for value, fmt in ((date(2026, 9, 12), "yyyy-mm-dd"), (123, "000000"), (12.5, "0.00"), (1234567890123456, "General")):
            def change(w, s):
                s["A2"] = value
                s["A2"].number_format = fmt
            with self.subTest(value=value), self.assertRaises(PortabilityError):
                self.parse(workbook_bytes(change=change))
        self.assertEqual(self.parse(workbook_bytes([["Date"], ["2026-09-12"]]))[3][0][1]["Date"], "2026-09-12")

    def test_empty_duplicate_numeric_and_missing_headers_rejected(self):
        for rows in ([["ID"]], [["ID", "ID"], ["a", "b"]], [[1], ["a"]], [["ID", None], ["a", "b"]]):
            with self.subTest(rows=rows), self.assertRaises(PortabilityError):
                self.parse(workbook_bytes(rows))

    def test_row_column_and_cell_bounds(self):
        for rows in ([["ID"]] + [[str(i)] for i in range(1001)],
                     [[str(i) for i in range(41)], ["value"] * 41],
                     [["ID"], ["x" * 4097]]):
            with self.subTest(size=len(rows)), self.assertRaises(PortabilityError):
                self.parse(workbook_bytes(rows))
        content = workbook_bytes([["ID"]] + [[str(i)] for i in range(1000)])
        self.assertEqual(len(self.parse(content)[3]), 1000)

    def test_false_dimensions_cannot_truncate_data(self):
        content = workbook_bytes([["ID", "Name"], ["a", "First"], ["b", "Second"]])
        content = rewrite(content, {"xl/worksheets/sheet1.xml": lambda d: d.replace(b'ref="A1:B3"', b'ref="A1:A1"')})
        self.assertEqual(len(self.parse(content)[3]), 2)
        for source, target in ((b'ref="A1:A1"', b'ref="A1:XFD1048576"'), (b'r="A3"', b'r="A1048576"'), (b'r="B3"', b'r="A3"')):
            with self.subTest(target=target), self.assertRaises(PortabilityError):
                self.parse(rewrite(content, {"xl/worksheets/sheet1.xml": lambda d: d.replace(source, target)}))

    def test_unsafe_duplicate_and_macro_zip_parts_rejected(self):
        for name in ("../escape.xml", "/absolute.xml", "C:/escape.xml", "xl\\bad.xml", "xl/vbaProject.bin", "xl/WORKBOOK.xml", "xl/workbook.xml"):
            with self.subTest(name=name), self.assertRaises(PortabilityError):
                self.parse(rewrite(workbook_bytes(), additions={name: b"<bad/>"}))
        info = zipfile.ZipInfo("docProps/custom.xml")
        info.external_attr = (stat.S_IFLNK | 0o777) << 16
        with self.assertRaises(PortabilityError):
            self.parse(rewrite(workbook_bytes(), additions={info: b"target"}))

    def test_zip_expansion_ratio_and_entry_limits(self):
        content = workbook_bytes()
        for constant, bound in (("MAX_EXPANDED", 100), ("MAX_PART", 100), ("MAX_MEMBERS", 2), ("MAX_RATIO", 1)):
            with self.subTest(constant=constant), patch.object(xlsx, constant, bound), self.assertRaises(PortabilityError):
                self.parse(content)

    def test_xml_entities_and_external_relationships_rejected(self):
        content = workbook_bytes()
        malicious = b'<!DOCTYPE worksheet [<!ENTITY x SYSTEM "file:///private">]><worksheet>&x;</worksheet>'
        with self.assertRaises(PortabilityError):
            self.parse(rewrite(content, {"xl/worksheets/sheet1.xml": lambda d: malicious}))
        for transform in (lambda d: d.replace(b'Target="/xl/worksheets/sheet1.xml"', b'Target="https://example.com/a" TargetMode="External"'),
                          lambda d: d.replace(b'Target="/xl/worksheets/sheet1.xml"', b'Target="../../escape.xml"')):
            with self.assertRaises(PortabilityError):
                self.parse(rewrite(content, {"xl/_rels/workbook.xml.rels": transform}))

    def test_bad_zip_wrong_extension_and_macro_content_type_rejected(self):
        for content, filename in ((b"not a zip", "bad.xlsx"), (workbook_bytes(), "bad.xls"), (workbook_bytes(), "bad.xlsm")):
            with self.assertRaises(PortabilityError):
                parse_source(content, filename)
        content = rewrite(workbook_bytes(), {"[Content_Types].xml": lambda d: d.replace(
            b"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml", b"application/vnd.ms-excel.sheet.macroEnabled.main+xml")})
        with self.assertRaises(PortabilityError):
            self.parse(content)

    def test_negative_shared_string_and_style_references_rejected(self):
        for transform in (lambda d: d.replace(b't="inlineStr"', b't="s"').replace(b"<is>", b"<v>-1</v><is>", 1),
                          lambda d: d.replace(b'r="A2"', b'r="A2" s="-1"')):
            with self.assertRaises(PortabilityError):
                self.parse(rewrite(workbook_bytes(), {"xl/worksheets/sheet1.xml": transform}))


@override_settings(STORAGES={"default": {"BACKEND": "django.core.files.storage.FileSystemStorage"},
                           "staticfiles": {"BACKEND": "django.contrib.staticfiles.storage.StaticFilesStorage"}})
class XlsxPortabilityTests(PortabilityFixture):
    def stage_xlsx(self, content=None, workspace=None, profile="party-master/1", system="paper"):
        return services.stage_import(workspace_id=(workspace or self.a).pk, actor=self.actor, content=content or workbook_bytes(),
            filename="register.xlsx", source_system=system, profile=profile)

    def save(self, batch):
        return presets.save_preset(workspace_id=self.a.pk, actor=self.actor, batch_id=batch.public_id,
            name="Register", approval_digest=batch.approval_digest)

    def test_master_mapping_commit_csv_replay_and_clean_workspace_roundtrip(self):
        with self.scoped():
            batch = self.stage_xlsx()
            self.assertEqual(batch.source_type, "xlsx")
            self.assertFalse(Party.objects.exists())
            batch = self.ready(batch)
            self.commit(batch)
            self.commit(self.ready())  # CSV and XLSX share exact source identity semantics.
            self.assertEqual(Party.objects.count(), 1)
            exported = services.export_parties(workspace_id=self.a.pk, actor=self.actor)
        with self.scoped(self.b):
            restored = self.stage(exported.content, exported.namespace, "master.jsonl", self.b)
            self.commit(self.ready(restored, {}, self.b), self.b)
            self.assertEqual(Party.objects.get().display_name, "Asha Devi")

    def test_csv_preset_applies_to_xlsx_and_xlsx_preset_to_csv(self):
        with self.scoped():
            csv_batch = self.ready()
            preset = self.save(csv_batch)
            batch = services.validate_import(workspace_id=self.a.pk, actor=self.actor,
                batch_id=self.stage_xlsx().public_id, preset_id=preset.public_id)
            self.assertEqual(batch.mapping_preset_id, preset.pk)
            self.assertEqual(self.save(batch).pk, preset.pk)
            self.commit(batch)
            mapping = copy.deepcopy(MAPPING)
            mapping["defaults"]["risk_label"] = "LOW"
            second = self.save(self.ready(self.stage_xlsx(workbook_bytes([["Legacy", "Party", "Phone"], ["new-2", "Bina Shah", ""]])), mapping))
            self.assertEqual(second.version, 2)
            csv_target = self.stage(b"Legacy,Party,Phone\nnew-2,Bina Shah,\n")
            applied = services.validate_import(workspace_id=self.a.pk, actor=self.actor, batch_id=csv_target.public_id, preset_id=second.public_id)
            self.commit(applied)
            self.assertEqual(Party.objects.count(), 2)

    def test_all_child_profiles_use_xlsx_and_preserve_roundtrip(self):
        from . import test_children, test_roles, test_relationships
        with self.scoped():
            self.commit(self.ready(self.stage_xlsx()))
            self.commit(self.ready(self.stage_xlsx(workbook_bytes([["Legacy", "Party", "Phone"], ["old-2", "Bina Shah", ""]]))))
            PartyRoleType.objects.create(key="BORROWER", label="Borrower")
            records = {
                contract.CONTACT: test_children.ChildPortabilityTests.record(self, contract.CONTACT),
                contract.ADDRESS: test_children.ChildPortabilityTests.record(self, contract.ADDRESS),
                contract.IDENTIFIER: {"party_source_system": "paper", "party_external_id": "old-1", "identifier_type": "PASSPORT", "value": "AB001234", "source_is_verified": False},
                contract.ROLE: test_roles.RolePortabilityTests.record(self),
                contract.RELATIONSHIP: test_relationships.RelationshipPortabilityTests.record(self),
            }
            exports = {}
            for profile, record in records.items():
                batch = self.stage_xlsx(workbook_bytes([["Legacy", *record], ["child-1", *record.values()]]), profile=profile, system="details")
                mapping = {"columns": {"Legacy": "source.external_id", **{k: k for k in record}}, "normalization": [
                    {"field": k, "rule": "boolean", "version": 1} for k, v in record.items() if type(v) is bool]}
                if profile == contract.ROLE:
                    mapping["role_type_map"] = {"customer": "BORROWER"}
                self.commit(self.ready(batch, mapping))
                exports[profile] = services.export_children(workspace_id=self.a.pk, actor=self.actor, profile=profile)
            master = services.export_parties(workspace_id=self.a.pk, actor=self.actor)
        with self.scoped(self.b):
            PartyRoleType.objects.create(key="BORROWER", label="Borrower")
            self.commit(self.ready(self.stage(master.content, master.namespace, "master.jsonl", self.b), {}, self.b), self.b)
            for profile, exported in exports.items():
                batch = services.stage_import(workspace_id=self.b.pk, actor=self.actor, profile=profile,
                    content=exported.content, source_system=exported.namespace, filename="child.jsonl")
                mapping = {"role_type_map": {"BORROWER": "BORROWER"}} if profile == contract.ROLE else {}
                self.commit(self.ready(batch, mapping, self.b), self.b)
                self.assertEqual(services.export_children(workspace_id=self.b.pk, actor=self.actor, profile=profile).count, 1)

    def test_malformed_workbook_leaves_no_batch_and_wrong_workspace_cannot_read(self):
        with self.scoped():
            with self.assertRaises(PortabilityError):
                self.stage_xlsx(b"bad zip")
            self.assertFalse(ImportBatch.objects.exists())
            batch = self.stage_xlsx()
        with self.scoped(self.b):
            self.assertFalse(ImportBatch.objects.exists())
            with self.assertRaises(PermissionDenied):
                services.preview_import(workspace_id=self.b.pk, actor=self.actor, batch_id=batch.public_id)

    def test_upload_mapping_and_preset_controls_render_for_xlsx(self):
        with self.scoped():
            self.save(self.ready())
            req = RequestFactory().post("/", {"profile": "party-master/1", "source_system": "paper",
                "source_file": SimpleUploadedFile("register.xlsx", workbook_bytes())})
            req.workspace, req.user = self.a, self.actor
            self.assertEqual(views.upload(req).status_code, 302)
            batch = ImportBatch.objects.get(source_type="xlsx")
            req = RequestFactory().get("/")
            req.workspace, req.user = self.a, self.actor
            response = views.batch_detail(req, batch.public_id)
            self.assertContains(response, "Saved mapping version")
            self.assertContains(response, "Map columns and validate")
            req = RequestFactory().post("/", {"action": "validate", "column_0": "source.external_id", "column_1": "name",
                "column_2": "primary_phone", "default_kind": "INDIVIDUAL", "default_status": "ACTIVE", "default_credit_hold": "false", "trim": "on"})
            req.workspace, req.user = self.a, self.actor
            self.assertEqual(views.batch_detail(req, batch.public_id).status_code, 302)
            batch.refresh_from_db()
            self.assertEqual(batch.state, "READY")

    def test_changed_xlsx_source_and_stale_approval_cannot_overwrite(self):
        with self.scoped():
            batch = self.ready(self.stage_xlsx())
            self.commit(batch)
            changed = self.ready(self.stage_xlsx(workbook_bytes([["Legacy", "Party", "Phone"], ["old-1", "Changed", ""]])))
            self.assertEqual(changed.summary["conflicts"], 1)
            with self.assertRaises(PortabilityError):
                self.commit(changed)
