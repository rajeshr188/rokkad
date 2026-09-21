"""Bounded, values-only XLSX input; no extraction, formulas or display-format guessing."""
import io
import posixpath
import re
import stat
import zipfile
from decimal import Decimal

from defusedxml.ElementTree import iterparse
from openpyxl import load_workbook
from openpyxl.utils.cell import coordinate_to_tuple, range_boundaries

from .parsers import MAX_CELL, MAX_COLUMNS, MAX_ROWS, PortabilityError, _check_limits

MAX_MEMBERS = 128
MAX_EXPANDED = 16 * 1024 * 1024
MAX_PART = 8 * 1024 * 1024
MAX_RATIO = 100
SHEET_NS = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
PARTS = re.compile(r"(?:\[Content_Types\]\.xml|_rels/\.rels|docProps/(?:app|core)\.xml|xl/(?:workbook\.xml|styles\.xml|sharedStrings\.xml|_rels/workbook\.xml\.rels|theme/theme[0-9]+\.xml|worksheets/sheet[0-9]+\.xml|worksheets/_rels/sheet[0-9]+\.xml\.rels))\Z")
REL_TYPES = {"officeDocument", "metadata/core-properties", "extended-properties", "worksheet", "styles", "theme", "sharedStrings"}
FORBIDDEN = {"f", "mergeCell", "hyperlink", "externalReference", "definedName", "autoFilter", "dataValidation",
             "conditionalFormatting", "extLst", "oleObject", "drawing", "legacyDrawing", "tablePart"}


def _xml(data):
    depth = count = 0
    parser = iterparse(io.BytesIO(data), events=("start", "end"), forbid_dtd=True, forbid_entities=True, forbid_external=True)
    for event, _ in parser:
        if event == "start":
            depth += 1
            count += 1
            if depth > 64 or count > 200000:
                raise PortabilityError("Workbook XML exceeds the structural limits.")
        else:
            depth -= 1
    return parser.root


def _relationships(root, name, names):
    seen = set()
    base = name.rsplit("/_rels/", 1)[0] if "/_rels/" in name else ""
    for rel in root:
        if rel.get("Id") in seen:
            raise PortabilityError("Duplicate workbook relationship identifiers are unsupported.")
        seen.add(rel.get("Id"))
        target = rel.get("Target", "")
        kind = rel.get("Type", "")
        if rel.get("TargetMode", "").lower() == "external" or not any(kind.endswith("/" + t) for t in REL_TYPES):
            raise PortabilityError("External links and embedded workbook features are unsupported.")
        if not target or any(c in target for c in ("\\", ":", "?", "#", "%")):
            raise PortabilityError("Invalid workbook relationship target.")
        resolved = posixpath.normpath(target.lstrip("/") if target.startswith("/") else posixpath.join(base, target))
        if resolved.startswith("../") or resolved not in names:
            raise PortabilityError("Workbook relationships must reference existing internal parts.")


def _worksheet(root):
    if root.tag != SHEET_NS + "worksheet":
        raise PortabilityError("Only standard XLSX worksheets are supported.")
    last_row = max_col = 0
    numeric = {}
    for node in root.iter():
        local = node.tag.rsplit("}", 1)[-1]
        if local in FORBIDDEN:
            raise PortabilityError("Use plain values without formulas, merged cells, filters, links or embedded features.")
        if local in {"row", "col", "sheetFormatPr"}:
            if node.get("hidden", "false").strip().lower() in {"1", "true"} or node.get("zeroHeight", "false").strip().lower() in {"1", "true"}:
                raise PortabilityError("Hidden rows or columns are unsupported.")
            if node.get("s", "0") != "0" or node.get("style", "0") != "0":
                raise PortabilityError("Whole-row and whole-column styles are unsupported; use plain cell values.")
            if any(node.get(k) is not None and float(node.get(k)) <= 0 for k in ("ht", "width", "defaultRowHeight")):
                raise PortabilityError("Zero-height rows or zero-width columns are unsupported.")
        if local == "dimension":
            bounds = range_boundaries(node.get("ref", ""))
            if any(v is None for v in bounds) or bounds[2] > MAX_COLUMNS or bounds[3] > MAX_ROWS + 1:
                raise PortabilityError("Worksheet dimensions exceed 40 columns or 1,001 physical rows including the header.")
    data = root.find(SHEET_NS + "sheetData")
    if data is None or len(root.findall(SHEET_NS + "sheetData")) != 1:
        raise PortabilityError("The worksheet contains no data.")
    if len(list(root.iter(SHEET_NS + "row"))) != len(data) or len(list(root.iter(SHEET_NS + "c"))) != sum(len(row) for row in data):
        raise PortabilityError("Worksheet cells must belong to a single data table.")
    for row in data:
        row_number = int(row.get("r", "0"))
        if row.tag != SHEET_NS + "row" or not last_row < row_number <= MAX_ROWS + 1:
            raise PortabilityError("Worksheet row numbers must be ordered and within the physical row limit.")
        last_row = row_number
        previous_col = 0
        for cell in row:
            coordinate = cell.get("r", "")
            if cell.tag != SHEET_NS + "c" or not re.fullmatch(r"[A-Z]{1,3}[1-9][0-9]{0,6}", coordinate):
                raise PortabilityError("Invalid worksheet cell coordinates.")
            cell_row, col = coordinate_to_tuple(coordinate)
            if cell_row != row_number or not previous_col < col <= MAX_COLUMNS:
                raise PortabilityError("Cells must have unique ordered coordinates within 40 columns.")
            previous_col, max_col = col, max(max_col, col)
            kind = cell.get("t", "n")
            if len(cell.findall(SHEET_NS + "v")) > 1 or len(cell.findall(SHEET_NS + "is")) > 1:
                raise PortabilityError("Duplicate cell values are unsupported.")
            value = cell.findtext(SHEET_NS + "v")
            if not re.fullmatch(r"[0-9]{1,6}", cell.get("s", "0")):
                raise PortabilityError("Invalid worksheet style reference.")
            if kind == "s" and (value is None or not re.fullmatch(r"[0-9]{1,6}", value)):
                raise PortabilityError("Invalid shared-string reference.")
            if kind not in {"n", "b", "s", "str", "inlineStr"}:
                raise PortabilityError("Error and date cells are unsupported; supply dates as ISO text.")
            if kind == "b" and value not in {"0", "1"}:
                raise PortabilityError("Invalid boolean cell.")
            if kind == "n" and value is not None:
                if len(value) > 128:
                    raise PortabilityError("Numeric cell exceeds the precision limit; use text.")
                number = Decimal(value)
                if not number.is_finite() or len(number.as_tuple().digits) > 15 or abs(number.adjusted()) > 308:
                    raise PortabilityError("Numeric cells must be finite with at most 15 significant digits; use text for identifiers.")
                numeric[coordinate] = format(number, "f")
    return last_row, max_col, numeric


def _preflight(content):
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        members = archive.infolist()
        if len(members) > MAX_MEMBERS or sum(p.file_size for p in members) > MAX_EXPANDED:
            raise PortabilityError("Workbook exceeds the ZIP entry or expanded-size limit.")
        names = set()
        for part in members:
            name = part.filename
            if (name.casefold() in names or name.startswith("/") or "\\" in name or ":" in name
                    or any(p in {"", ".", ".."} for p in name.split("/"))
                    or stat.S_ISLNK(part.external_attr >> 16) or part.flag_bits & 1):
                raise PortabilityError("Unsafe, duplicate or encrypted ZIP members are unsupported.")
            names.add(name.casefold())
            if not PARTS.fullmatch(name):
                raise PortabilityError("Use a plain XLSX data workbook without macros, attachments or unsupported package parts.")
            if part.compress_type not in {zipfile.ZIP_STORED, zipfile.ZIP_DEFLATED} or part.file_size > MAX_PART or part.file_size > max(part.compress_size, 1) * MAX_RATIO:
                raise PortabilityError("Workbook exceeds a ZIP member size or compression-ratio limit.")
        exact_names = {p.filename for p in members}
        required = {"[Content_Types].xml", "_rels/.rels", "xl/workbook.xml", "xl/_rels/workbook.xml.rels"}
        if not required <= exact_names:
            raise PortabilityError("The file is not a supported XLSX workbook.")
        sheets = []
        for part in members:
            with archive.open(part) as stream:
                data = stream.read(MAX_PART + 1)
            if len(data) != part.file_size or len(data) > MAX_PART:
                raise PortabilityError("Invalid workbook part size.")
            root = _xml(data)
            if part.filename.endswith(".rels"):
                _relationships(root, part.filename, exact_names)
            elif part.filename == "[Content_Types].xml":
                workbook_types = [n.get("ContentType") for n in root if n.get("PartName") == "/xl/workbook.xml"]
                if workbook_types != ["application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"]:
                    raise PortabilityError("Only non-macro XLSX workbooks are supported.")
            elif part.filename == "xl/workbook.xml":
                for node in root.iter():
                    if node.tag.rsplit("}", 1)[-1] in FORBIDDEN:
                        raise PortabilityError("Defined names and external workbook features are unsupported.")
            if re.fullmatch(r"xl/worksheets/sheet[0-9]+\.xml", part.filename):
                sheets.append(_worksheet(root))
        if len(sheets) != 1:
            raise PortabilityError("Use exactly one visible worksheet per upload.")
        return sheets[0]


def parse_xlsx(content):
    workbook = None
    try:
        max_row, max_col, numeric = _preflight(content)
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=False, keep_links=False)
        if len(workbook.sheetnames) != 1 or len(workbook.worksheets) != 1 or workbook.worksheets[0].sheet_state != "visible":
            raise PortabilityError("Use exactly one visible worksheet per upload.")
        sheet = workbook.worksheets[0]
        sheet.reset_dimensions()  # Actual checked coordinates, not untrusted dimension hints.
        rows, headers = [], []
        for number, cells in enumerate(sheet.iter_rows(min_row=1, max_row=max_row, max_col=max_col), 1):
            values = []
            for cell in cells:
                value = cell.value
                if value is None:
                    values.append("")
                    continue
                if cell.data_type == "b":
                    value = "true" if value else "false"
                elif cell.data_type == "n" and cell.number_format == "General":
                    value = numeric[cell.coordinate]
                elif cell.data_type != "s":
                    raise PortabilityError(f"Cell {cell.coordinate} requires plain text; Excel dates and formatted numbers are unsupported.")
                if not isinstance(value, str) or "\x00" in value or len(value) > MAX_CELL:
                    raise PortabilityError("A worksheet cell is invalid or exceeds 4,096 characters.")
                if number == 1 and cell.data_type != "s":
                    raise PortabilityError("Worksheet headers must be text in row 1.")
                values.append(value)
            if number == 1:
                headers = values
                if not headers or len(set(headers)) != len(headers) or any(not h.strip() or len(h) > 255 for h in headers):
                    raise PortabilityError("Use 1–40 unique, nonempty text headers in worksheet row 1.")
            elif any(v != "" for v in values):
                rows.append((number, dict(zip(headers, values))))
                _check_limits(rows)
        if not rows:
            raise PortabilityError("The worksheet contains no records.")
        return headers, rows
    except PortabilityError:
        raise
    except Exception as exc:
        # Malformed ZIP/XML/style/string references have many library exception
        # types. Never expose package contents or arbitrary parser text to users.
        raise PortabilityError("Malformed or unsupported XLSX workbook. Use a plain values-only worksheet.") from exc
    finally:
        if workbook is not None:
            workbook.close()
