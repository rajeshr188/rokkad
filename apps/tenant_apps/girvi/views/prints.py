import logging

from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.db.models import CharField, Count, OuterRef, Q, Subquery, Sum, Value
from django.db.models.functions import Coalesce
from django.http import HttpResponse
from django.shortcuts import get_object_or_404, redirect, render
from openpyxl import Workbook
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import inch, mm
from reportlab.pdfgen import canvas
from reportlab.platypus import (
    BalancedColumns,
    ListFlowable,
    ListItem,
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Table,
    TableStyle,
)

from apps.tenant_apps.girvi.documents.loan_ticket import (
    get_custom_jcl,
    grid_template,
    print_labels_pdf,
)
from apps.tenant_apps.girvi.models.template import LoanTemplate
from apps.tenant_apps.girvi.service_modules.printing import LoanPrintService
from apps.tenant_apps.girvi.service_modules.print_selection import (
    unreleased_given_loan_selection,
)
from apps.tenant_apps.notify.models import Notification
from apps.tenant_apps.notify.services import (
    DEFAULT_LOAN_REMINDER_CODE,
    create_bulk_loan_reminder_group,
)
from apps.tenant_apps.notify_v2.models import NotificationJob
from apps.tenant_apps.notify_v2.services import create_girvi_reminder_batch

from ..forms import LoanSelectionForm

logger = logging.getLogger(__name__)

_NOTIFY_V2_EVENT_MAP = {
    "LOAN_FIRST_REMINDER": "loan.first_reminder_due",
    "LOAN_SECOND_REMINDER": "loan.second_reminder_due",
    "LOAN_FINAL_NOTICE": "loan.final_notice_due",
    "LOAN_AUCTION_NOTICE": "loan.auction_notice_due",
}

_NOTIFY_V2_CHANNEL_MAP = {
    Notification.MediumType.Email: NotificationJob.Channel.EMAIL,
    Notification.MediumType.Letter: NotificationJob.Channel.LETTER,
    Notification.MediumType.Post: NotificationJob.Channel.POST,
    Notification.MediumType.SMS: NotificationJob.Channel.SMS,
    Notification.MediumType.Whatsapp: NotificationJob.Channel.WHATSAPP,
}


def print_labels(request):
    selection = unreleased_given_loan_selection(
        post_data=request.POST,
        query_data=request.GET,
    )
    if not selection.is_valid:
        if selection.invalid_ids:
            logger.warning("print_labels rejected invalid IDs: %s", selection.invalid_ids)
        return HttpResponse(status=400, content=selection.error)

    if selection.loans:
        form = LoanSelectionForm(initial={"loans": selection.loans})
        template_name = "girvi/loan/print_labels.html#content" if request.htmx else "girvi/loan/print_labels.html"
        return render(request, template_name, {"form": form})

    return HttpResponse(status=200, content="No unreleased loans selected.")


def print_label(request):
    if request.method == "POST":
        form = LoanSelectionForm(request.POST)
        if form.is_valid():
            loans = form.cleaned_data["loans"]
            return print_labels_pdf(loans)

        template_name = "girvi/loan/print_labels.html#content" if request.htmx else "girvi/loan/print_labels.html"
        return render(request, template_name, {"form": form})

    else:
        form = LoanSelectionForm()
        template_name = "girvi/loan/print_labels.html#content" if request.htmx else "girvi/loan/print_labels.html"
        return render(request, template_name, {"form": form})


@login_required
def notify_print(request):
    selection = unreleased_given_loan_selection(
        post_data=request.POST,
        query_data=request.GET,
    )
    if not selection.is_valid:
        if selection.invalid_ids:
            logger.warning("notify_print rejected invalid IDs: %s", selection.invalid_ids)
        return HttpResponse(status=400, content=selection.error)

    if selection.loans:
        notice_code = request.POST.get("notice_code", DEFAULT_LOAN_REMINDER_CODE)
        medium_type = request.POST.get(
            "medium_type",
            Notification.MediumType.Letter,
        )

        try:
            ng = create_bulk_loan_reminder_group(
                selection.loans,
                notice_code=notice_code,
                medium_type=medium_type,
            )
        except ValueError as exc:
            return HttpResponse(status=400, content=str(exc))
        except Exception:
            logger.exception("Error creating reminder notifications.")
            return HttpResponse(status=500, content="Error creating notifications.")

        return redirect(ng.get_absolute_url())

    return HttpResponse(status=200, content="No unreleased loans selected.")


@login_required
def notify_print_v2(request):
    selection = unreleased_given_loan_selection(
        post_data=request.POST,
        query_data=request.GET,
    )
    if not selection.is_valid:
        if selection.invalid_ids:
            logger.warning("notify_print_v2 rejected invalid IDs: %s", selection.invalid_ids)
        return HttpResponse(status=400, content=selection.error)

    selected_loans = list(selection.loans or [])
    if not selected_loans:
        return HttpResponse(status=200, content="No unreleased loans selected.")

    notice_code = request.POST.get("notice_code", DEFAULT_LOAN_REMINDER_CODE)
    medium_type = request.POST.get("medium_type", Notification.MediumType.Letter)
    event_key = _NOTIFY_V2_EVENT_MAP.get(notice_code, "loan.first_reminder_due")
    channel = _NOTIFY_V2_CHANNEL_MAP.get(medium_type, NotificationJob.Channel.LETTER)

    try:
        batch_result = create_girvi_reminder_batch(
            loans=selected_loans,
            created_by=request.user,
            event_key=event_key,
            channel=channel,
        )
    except ValueError as exc:
        return HttpResponse(status=400, content=str(exc))
    except Exception:
        logger.exception("Error creating notify_v2 reminder batch.")
        return HttpResponse(status=500, content="Error creating notify_v2 batch.")

    try:
        messages.success(
            request,
            f"Created notify_v2 batch with {batch_result.preview.borrower_count} recipient(s).",
        )
    except Exception:
        pass
    destination = batch_result.batch.get_absolute_url()
    if request.headers.get("HX-Request") == "true":
        # HTMX requests from bulk-action dropdown should perform a full navigation.
        return HttpResponse(status=204, headers={"HX-Redirect": destination})
    return redirect(destination)


import base64


@login_required
def print_loan(request, pk=None):
    loan = get_object_or_404(
        GivenLoan.objects.select_related(
            "borrower", "series", "series__license"
        ).prefetch_related(
            "loanitems", "borrower__address", "borrower__contactno"
        ),
        pk=pk,
    )
    result = LoanPrintService.build_print_result(
        loan,
        template_resolver=LoanTemplate.objects.get_default,
        renderer=get_custom_jcl,
    )
    if not result.template:
        messages.warning(request, result.error_message)
        return redirect("girvi:girvi_loan_detail", pk=loan.pk)

    if not result.ok:
        readiness = result.readiness or {}
        logger.error(
            "Loan PDF generation failed for loan=%s template_id=%s print_option=%s readiness=%s",
            getattr(loan, "loan_id", None),
            getattr(result.template, "pk", None),
            getattr(result.template, "print_option", None),
            readiness.get("summary"),
        )
        messages.error(request, result.error_message)
        return redirect("girvi:girvi_loan_detail", pk=loan.pk)

    response = HttpResponse(result.pdf, content_type="application/pdf")
    response["Content-Disposition"] = f"inline; filename='{loan.loan_id}.pdf'"
    response["Content-Transfer-Encoding"] = "binary"
    return response
    # Encode the PDF in base64
    # pdf_base64 = base64.b64encode(pdf).decode("utf-8")

    # # Render the object HTML
    # object_html = f"""
    # <object data="data:application/pdf;base64,{pdf_base64}" type="application/pdf" width="100%" height="600px">
    #     <p>Your browser does not support PDFs. <a href="data:application/pdf;base64,{pdf_base64}">Download the PDF</a>.</p>
    # </object>
    # """
    # return HttpResponse(object_html)


@login_required
def print_grid_template(request):
    pdf = grid_template()

    # Encode the PDF in base64
    pdf_base64 = base64.b64encode(pdf).decode("utf-8")

    # Render the object HTML
    object_html = f"""
    <object data="data:application/pdf;base64,{pdf_base64}" type="application/pdf" width="100%" height="600px">
        <p>Your browser does not support PDFs. <a href="data:application/pdf;base64,{pdf_base64}">Download the PDF</a>.</p>
    </object>
    """
    return HttpResponse(object_html)


class PageNumCanvas(canvas.Canvas):
    def __init__(self, *args, **kwargs):
        canvas.Canvas.__init__(self, *args, **kwargs)
        self.pages = []

    def showPage(self):
        self.pages.append(dict(self.__dict__))
        self._startPage()

    def save(self):
        page_count = len(self.pages)
        for page in self.pages:
            self.__dict__.update(page)
            self.draw_header_footer(page_count)
            canvas.Canvas.showPage(self)
        canvas.Canvas.save(self)

    def draw_header_footer(self, page_count):
        self.setFont("Helvetica", 10)
        self.drawRightString(
            200 * mm, 20 * mm, "Page %d of %d" % (self._pageNumber, page_count)
        )

        # Draw header
        self.saveState()
        styles = getSampleStyleSheet()
        header = Paragraph("Loan Report", styles["Heading1"])
        w, h = header.wrap(self._pagesize[0] - 1 * inch, self._pagesize[1])
        header.drawOn(self, inch, self._pagesize[1] - inch - h + 20)
        self.restoreState()


# @login_required
# def generate_loans_ledger_pdf(response):
#     # add annotation,bookmarks,table of contents,pagination,header and footer
#     response = HttpResponse(content_type="application/pdf")
#     response["Content-Disposition"] = 'attachment; filename="loan_report.pdf"'

#     doc = SimpleDocTemplate(response, pagesize=landscape(A4))
#     elements = []

#     styles = getSampleStyleSheet()
#     cover_title = Paragraph("Loan Report", styles["Title"])
#     cover_subtitle = Paragraph("Generated on: [Date]", styles["Heading2"])
#     elements.append(Spacer(1, 2 * inch))
#     elements.append(cover_title)
#     elements.append(Spacer(1, 0.5 * inch))
#     elements.append(cover_subtitle)
#     elements.append(PageBreak())

#     # Add Table of Contents
#     toc = TableOfContents()
#     toc.levelStyles = [
#         ParagraphStyle(fontSize=14, name="Heading1", leading=16),
#         ParagraphStyle(fontSize=12, name="Heading2", leading=14),
#         ParagraphStyle(fontSize=10, name="Heading3", leading=12),
#     ]
#     elements.append(Paragraph("Table of Contents", getSampleStyleSheet()["Heading1"]))
#     elements.append(toc)
#     elements.append(PageBreak())

#     # Add Table
#     data = [
#         [
#             "Loan ID",
#             "Loan Date",
#             "Customer",
#             "Loan Amount",
#             "Weight",
#             "Present Value",
#             "Item Description",
#             "Release Date",
#             "Released By",
#         ]
#     ]
#     loans = Loan.objects.filter(series__is_active = True).select_related("customer", "release")[:500]
#     for loan in loans.iterator(chunk_size=100):
#         data.append(
#             [
#                 loan.loan_id,
#                 loan.loan_date.date(),
#                 loan.customer,
#                 loan.loan_amount,
#                 loan.weight,
#                 loan.value,
#                 loan.item_desc,
#                 loan.release.release_date.date() if loan.is_released else "",
#                 loan.release.released_by if loan.is_released else "",
#             ]
#         )
#     # Define column widths
#     col_widths = [
#         0.5 * inch,
#         1 * inch,
#         1 * inch,
#         1 * inch,
#         1 * inch,
#         1 * inch,
#         2.5 * inch,
#         1 * inch,
#         1 * inch,
#     ]
#     table = Table(data, colWidths=col_widths,repeatRows=1, splitByRow=1)
#     table.setStyle(
#         TableStyle(
#             [
#                 ("BACKGROUND", (0, 0), (-1, 0), colors.grey),
#                 ("TEXTCOLOR", (0, 0), (-1, 0), colors.whitesmoke),
#                 ("ALIGN", (0, 0), (-1, -1), "CENTER"),
#                 ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
#                 ("BOTTOMPADDING", (0, 0), (-1, 0), 12),
#                 ("BACKGROUND", (0, 1), (-1, -1), colors.beige),
#                 ("GRID", (0, 0), (-1, -1), 1, colors.black),
#             ]
#         )
#     )

#     elements.append(table)

#     # Define a frame for the content
#     frame = Frame(
#         inch,
#         inch,
#         landscape(A4)[0] - 2 * inch,
#         landscape(A4)[1] - 2 * inch,
#         id="normal",
#     )

#     # Create a PageTemplate with the frame
#     template = PageTemplate(id="test", frames=frame, onPage=PageNumCanvas)

#     # Build PDF with the template
#     doc.addPageTemplates([template])
#     # Build PDF
#     doc.multiBuild(elements, canvasmaker=PageNumCanvas)

#     return response


# @login_required
# def generate_loans_ledger_pdf(request):
#     # Create the HttpResponse object with the appropriate PDF headers.
#     response = HttpResponse(content_type='application/pdf')
#     response['Content-Disposition'] = 'attachment; filename="loans_ledger.pdf"'

#     # Create the PDF object, using the response object as its "file."
#     doc = SimpleDocTemplate(response, pagesize=letter)
#     styles = getSampleStyleSheet()

#     # Create a list to hold the elements for the PDF.
#     elements = []

#     # Add the header
#     header = Paragraph("Loan Report", styles["Heading1"])
#     elements.append(header)

#     # Query the Series model to get all active series
#     series_list = Series.objects.filter(is_active=True)

#     for series in series_list:
#         # Add series title
#         series_title = Paragraph(f"Series: {series.name}", styles["Heading2"])
#         elements.append(series_title)

#         # Add Table
#         data = [
#             [
#                 "Loan ID",
#                 "Loan Date",
#                 "Customer",
#                 "Loan Amount",
#                 "Weight",
#                 "Present Value",
#                 "Item Description",
#                 "Release Date",
#                 "Released By",
#             ]
#         ]
#         loans = Loan.objects.filter(series=series).select_related("customer", "release")[:50]
#         for loan in loans.iterator(chunk_size=100):
#             data.append(
#                 [
#                     loan.loan_id,
#                     loan.loan_date.date(),
#                     loan.customer,
#                     loan.loan_amount,
#                     loan.weight,
#                     loan.value,
#                     loan.item_desc,
#                     loan.release.release_date.date() if loan.is_released else "",
#                     loan.release.released_by if loan.is_released else "",
#                 ]
#             )
#         # Define column widths
#         col_widths = [
#             0.5 * inch,
#             1 * inch,
#             1 * inch,
#             1 * inch,
#             1 * inch,
#             1 * inch,
#             2.5 * inch,
#             1 * inch,
#             1 * inch,
#         ]
#         table = Table(data, colWidths=col_widths, repeatRows=1, splitByRow=1)
#         table.setStyle(
#             TableStyle(
#                 [
#                     ("BACKGROUND", (0, 0), (-1, 0), colors.grey),
#                     ("TEXTCOLOR", (0, 0), (-1, 0), colors.whitesmoke),
#                     ("ALIGN", (0, 0), (-1, -1), "CENTER"),
#                     ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
#                     ("BOTTOMPADDING", (0, 0), (-1, 0), 12),
#                     ("BACKGROUND", (0, 1), (-1, -1), colors.beige),
#                     ("GRID", (0, 0), (-1, -1), 1, colors.black),
#                 ]
#             )
#         )
#         elements.append(table)

#     # Build the PDF
#     doc.build(elements)

#     return response


@login_required
def generate_unreleased_pdf(request):
    numbers = GivenLoan.objects.unreleased().values_list("loan_id", flat=True)

    response = HttpResponse(content_type="application/pdf")
    response["Content-Disposition"] = 'attachment; filename="numbers_report.pdf"'

    doc = SimpleDocTemplate(response, pagesize=A4)
    elements = []

    list_style = getSampleStyleSheet()["OrderedList"]
    list_style.bulletFontName = "Courier"
    list_style.bulletFontSize = 10
    list_style.leftIndent = 24
    list_items_builtins = [
        ListItem(Paragraph(b, getSampleStyleSheet()["Normal"])) for b in numbers
    ]
    elements.append(
        BalancedColumns(
            F=[ListFlowable(list_items_builtins, style=list_style)], nCols=3
        )
    )
    elements.append(PageBreak())

    doc.build(elements)
    return response

from ..models import GivenLoan, Series
from ..resources import LedgerResource


@login_required
def export_loans_to_excel(request):
    # Create a workbook
    wb = Workbook()
    wb.remove(wb.active)  # Remove the default sheet

    # Fetch all series
    series_list = Series.objects.filter(is_active=True)

    for series in series_list:
        # Create a new worksheet for each series
        ws = wb.create_sheet(title=series.name)

        # Get loans for the current series
        loans = GivenLoan.objects.filter(series=series).for_table_display()

        # Use LoanResource to export data
        loan_resource = LedgerResource()
        dataset = loan_resource.export(loans)

        # Add headers to the worksheet
        headers = dataset.headers
        ws.append(headers)

        # Add loan data to the worksheet
        for row in dataset:
            ws.append(row)

    # Save the workbook to a response
    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response[
        "Content-Disposition"
    ] = f'attachment; filename="{request.user.profile.workspace.name}_ledger.xlsx"'
    wb.save(response)

    return response


def _inventory_audit_queryset(scope="all", from_date=None, to_date=None):
    from apps.tenant_apps.contact.models import Address

    default_address_qs = Address.objects.filter(customer_id=OuterRef("borrower_id")).order_by(
        "-is_default", "-created", "-pk"
    )

    loans = (
        GivenLoan.objects.select_related("series", "borrower", "release")
        .prefetch_related("loanitems")
        .annotate(
            borrower_door=Coalesce(
                Subquery(default_address_qs.values("door_number")[:1]),
                Value(""),
                output_field=CharField(),
            ),
            borrower_street=Coalesce(
                Subquery(default_address_qs.values("street")[:1]),
                Value(""),
                output_field=CharField(),
            ),
            borrower_area=Coalesce(
                Subquery(default_address_qs.values("area")[:1]),
                Value(""),
                output_field=CharField(),
            ),
            borrower_city=Coalesce(
                Subquery(default_address_qs.values("city")[:1]),
                Value(""),
                output_field=CharField(),
            ),
            borrower_state=Coalesce(
                Subquery(default_address_qs.values("state")[:1]),
                Value(""),
                output_field=CharField(),
            ),
            borrower_zip=Coalesce(
                Subquery(default_address_qs.values("zip_code")[:1]),
                Value(""),
                output_field=CharField(),
            ),
            item_count=Count("loanitems", distinct=True),
            principal_amount=Sum("loanitems__loanamount"),
            total_weight=Sum("loanitems__weight"),
            in_vault=Count(
                "loanitems",
                filter=Q(loanitems__custody_status="in_vault"),
                distinct=True,
            ),
            with_lender=Count(
                "loanitems",
                filter=Q(loanitems__custody_status="with_lender"),
                distinct=True,
            ),
            with_customer=Count(
                "loanitems",
                filter=Q(loanitems__custody_status="with_customer"),
                distinct=True,
            ),
        )
        .order_by("series__name", "loan_id")
    )

    if scope == "unreleased":
        loans = loans.filter(release__isnull=True)
    elif scope == "released":
        loans = loans.filter(release__isnull=False)

    # Apply date range filter if provided
    if from_date:
        loans = loans.filter(loan_date__gte=from_date)
    if to_date:
        loans = loans.filter(loan_date__lte=to_date)

    return loans


def _build_inventory_audit_rows(*, scope="all", from_date=None, to_date=None):
    loans = _inventory_audit_queryset(scope, from_date=from_date, to_date=to_date)

    rows = []
    for loan in loans:
        # Build item details from prefetched items (no additional DB hits)
        items = list(loan.loanitems.all())
        item_details = []
        for item in items:
            detail = f"{item.itemtype}"
            if item.itemdesc:
                detail += f" ({item.itemdesc})"
            if item.weight:
                detail += f" {item.weight}g"
            item_details.append(detail)

        items_description = " | ".join(item_details) if item_details else "N/A"

        borrower_name = str(getattr(loan, "borrower", ""))
        address_parts = [
            loan.borrower_door,
            loan.borrower_street,
            loan.borrower_area,
            loan.borrower_city,
            loan.borrower_state,
            loan.borrower_zip,
        ]
        borrower_address = ", ".join(part for part in address_parts if part)

        item_count = loan.item_count or 0
        in_vault = loan.in_vault or 0
        with_lender = loan.with_lender or 0
        with_customer = loan.with_customer or 0
        physical_expected = in_vault + with_lender

        is_released = hasattr(loan, "release")

        # For unreleased loans, customer custody is a strong signal to review.
        if not is_released and with_customer > 0:
            audit_flag = "REVIEW"
            audit_note = "Unreleased loan has item(s) with customer"
        elif physical_expected != item_count:
            audit_flag = "REVIEW"
            audit_note = "Custody totals do not match item count"
        else:
            audit_flag = "OK"
            audit_note = "In sync"

        rows.append(
            {
                "series": getattr(getattr(loan, "series", None), "name", ""),
                "loan_id": loan.loan_id,
                "loan_date": loan.loan_date,
                "borrower": borrower_name,
                "borrower_address": borrower_address,
                "pawner_full": (
                    f"{borrower_name}, {borrower_address}"
                    if borrower_address
                    else borrower_name
                ),
                "principal_amount": loan.principal_amount or 0,
                "total_weight": loan.total_weight or 0,
                "items_description": items_description,
                "status": loan.status,
                "release_id": getattr(getattr(loan, "release", None), "release_id", ""),
                "release_date": getattr(getattr(loan, "release", None), "release_date", None),
                "item_count": item_count,
                "in_vault": in_vault,
                "with_lender": with_lender,
                "with_customer": with_customer,
                "physical_expected": physical_expected,
                "audit_flag": audit_flag,
                "audit_note": audit_note,
            }
        )
    return rows


@login_required
def export_active_loans_inventory_audit(request):
    from datetime import datetime
    
    export_format = (request.GET.get("format") or "xlsx").strip().lower()
    scope = (request.GET.get("scope") or "all").strip().lower()
    if scope not in {"all", "released", "unreleased"}:
        scope = "all"

    # Parse date range parameters
    from_date = None
    to_date = None
    try:
        if request.GET.get("from_date"):
            from_date = datetime.strptime(request.GET.get("from_date"), "%Y-%m-%d").date()
        if request.GET.get("to_date"):
            to_date = datetime.strptime(request.GET.get("to_date"), "%Y-%m-%d").date()
    except (ValueError, TypeError):
        pass

    rows = _build_inventory_audit_rows(scope=scope, from_date=from_date, to_date=to_date)

    scope_title_map = {
        "all": "Released + Unreleased Given Loans",
        "released": "Released Given Loans",
        "unreleased": "Unreleased Given Loans",
    }
    scope_filename_map = {
        "all": "all_loans_inventory_audit",
        "released": "released_loans_inventory_audit",
        "unreleased": "unreleased_loans_inventory_audit",
    }
    scope_title = scope_title_map[scope]
    scope_filename = scope_filename_map[scope]

    headers = [
        "Series",
        "Loan ID",
        "Loan Date",
        "Borrower",
        "Status",
        "Release ID",
        "Release Date",
        "Item Count",
        "In Vault",
        "With Lender",
        "With Customer",
        "Expected Physical Count",
        "Audit Flag",
        "Audit Note",
    ]

    if export_format == "pdf":
        response = HttpResponse(content_type="application/pdf")
        response["Content-Disposition"] = (
            f'attachment; filename="{scope_filename}.pdf"'
        )

        # Pledgebook-style renderer with fixed columns and fast canvas drawing.
        pdf = canvas.Canvas(response, pagesize=landscape(A4))
        width, height = landscape(A4)
        left = 16
        right = width - 16
        top = height - 16

        col_widths = [52, 130, 52, 58, 64, 42, 58, 116, 52, 86, 86]
        x_positions = [left]
        for w in col_widths:
            x_positions.append(x_positions[-1] + w)

        def _truncate(value, max_len):
            text = "" if value is None else str(value)
            return text if len(text) <= max_len else f"{text[:max_len-1]}…"

        def _draw_title_block():
            # Title block is 72pt tall to accommodate multi-line notes on the right
            block_h = 72
            pdf.setStrokeColor(colors.black)
            pdf.setLineWidth(0.8)
            pdf.rect(left, top - block_h, right - left, block_h, stroke=1, fill=0)

            # Centre title
            mid_x = (left + right) / 2
            pdf.setFont("Helvetica-Bold", 13)
            pdf.drawCentredString(mid_x, top - 18, "FORM E  Pledge Book")
            pdf.setFont("Helvetica", 8)
            pdf.drawCentredString(mid_x, top - 30, "(Section 10(1) (a) & Rule 7)")

            # Right-side notes column
            notes_x = right - 220
            note_lines = [
                "Rules framed under the Chennai Pawn Brokers Act",
                "Note: All entries in the pledge book except items 5, 9 and 11",
                "respecting each pledge shall be made on the day of pawning thereof.",
                "(4) Rate of interest charged: 16%",
                "(5) The time agreed upon for the redemption of pawn is 12 months.",
            ]
            pdf.setFont("Helvetica", 6.5)
            for idx, line in enumerate(note_lines):
                pdf.drawString(notes_x, top - 12 - (idx * 10), line)

            # Scope label bottom-left
            pdf.setFont("Helvetica", 7.5)
            pdf.drawString(left + 6, top - block_h + 8, f"Scope: {scope_title}")


        def _draw_table_header(y):
            headers_pledgebook = [
                "No. of\nPledge",
                "Name of pawner\nand full address",
                "Date of\nLoan",
                "Amount of\nPrincipal\nRs.",
                "Amount of every\npayment received\nRs.",
                "Weight\n(g)",
                "Present\nValue\nRs.",
                "Full and detailed\ndescription of\narticles",
                "Date of\nrelease",
                "Name and address\nof owner if\nnot redeemed",
                "Name and address\nof person\nredeeming",
            ]
            row_h = 32
            for i, text in enumerate(headers_pledgebook):
                x0 = x_positions[i]
                cw = col_widths[i]
                pdf.rect(x0, y - row_h, cw, row_h, stroke=1, fill=0)
                pdf.setFont("Helvetica", 6.7)
                lines = text.split("\n")
                for li, line in enumerate(lines):
                    pdf.drawString(x0 + 2, y - 9 - (li * 8), line)
            return y - row_h

        def _draw_data_row(y, row):
            row_h = 14
            values = [
                _truncate(row["loan_id"], 11),
                _truncate(row["pawner_full"], 42),
                row["loan_date"].strftime("%d/%m/%y") if row["loan_date"] else "",
                str(row["principal_amount"]),
                "-",
                _truncate(f"{row['total_weight']}", 8),
                str(row["principal_amount"]),
                _truncate(row["items_description"], 38),
                row["release_date"].strftime("%d/%m/%y") if row["release_date"] else "",
                _truncate(row["borrower"] if not row["release_date"] else "-", 28),
                _truncate(row["borrower"] if row["release_date"] else "-", 28),
            ]

            pdf.setFont("Helvetica", 6.8)
            for i, value in enumerate(values):
                x0 = x_positions[i]
                cw = col_widths[i]
                pdf.rect(x0, y - row_h, cw, row_h, stroke=1, fill=0)
                pdf.drawString(x0 + 2, y - 10, _truncate(value, max(5, int(cw / 4))))
            return y - row_h

        _draw_title_block()
        y = top - 76
        y = _draw_table_header(y)

        prev_series = None
        for row in rows:
            # Start new page when series changes
            if prev_series is not None and prev_series != row["series"]:
                pdf.showPage()
                _draw_title_block()
                y = top - 76
                y = _draw_table_header(y)
            
            if y < 24:
                pdf.showPage()
                _draw_title_block()
                y = top - 76
                y = _draw_table_header(y)

            y = _draw_data_row(y, row)
            prev_series = row["series"]

        pdf.save()
        return response

    # Default: xlsx
    wb = Workbook(write_only=True)
    ws = wb.create_sheet("Loan Inventory")
    ws.append(headers)

    for row in rows:
        ws.append(
            [
                row["series"],
                row["loan_id"],
                row["loan_date"].strftime("%Y-%m-%d") if row["loan_date"] else "",
                row["borrower"],
                row["status"],
                row["release_id"],
                row["release_date"].strftime("%Y-%m-%d") if row["release_date"] else "",
                row["item_count"],
                row["in_vault"],
                row["with_lender"],
                row["with_customer"],
                row["physical_expected"],
                row["audit_flag"],
                row["audit_note"],
            ]
        )

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = (
        f'attachment; filename="{scope_filename}.xlsx"'
    )
    wb.save(response)
    return response
