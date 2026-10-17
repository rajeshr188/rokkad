"""Excel presentation of the same complete, filtered loan directory."""

from datetime import date
from decimal import Decimal
from io import BytesIO

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from apps.tenant_apps.loans.domain.vocabulary import CollateralMetal
from apps.tenant_apps.loans.selectors.directory import loan_directory_export_rows
from apps.tenant_apps.loans.selectors.loan_summary import (
    get_loan_collateral_summary,
    get_loan_summary,
    completed_months,
    MONTHS_BASIS,
)

COLUMNS = (
    "Record type",
    "Loan",
    "Customer",
    "Customer code",
    "Licence",
    "Series",
    "Status",
    "Loan date",
    "Loan age (completed months)",
    "Months with unpaid interest",
    "Principal outstanding (INR)",
    "Interest pending (INR)",
    "Fees pending (INR)",
    "Total outstanding (INR)",
    "Balance basis",
    "Unpaid months basis",
    "Collateral summary",
    *(
        f"Held {metal.value.lower()} {weight} weight (g)"
        for metal in CollateralMetal
        for weight in ("gross", "net")
    ),
)


def _append(sheet, values, *, header=False, weight_columns_from=None):
    cells = []
    for index, value in enumerate(values):
        cell = WriteOnlyCell(sheet, value=value)
        # Customer names, original references and filter text are data, including
        # strings beginning with '='. Never emit customer-controlled formulas.
        if isinstance(value, str):
            cell.data_type = "s"
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        elif isinstance(value, Decimal):
            cell.number_format = (
                "#,##0.0000"
                if weight_columns_from is not None and index >= weight_columns_from
                else "#,##0.00"
            )
        elif isinstance(value, date):
            cell.number_format = "dd/mm/yyyy"
        if header:
            cell.font = Font(bold=True)
        cells.append(cell)
    sheet.append(cells)


def export_loan_directory(*, workspace, loan_filter, as_of_date):
    workbook = Workbook(write_only=True)
    summary_sheet = workbook.create_sheet("Summary")
    sheet = workbook.create_sheet("Loans")
    sheet.freeze_panes = "C2"
    for index, label in enumerate(COLUMNS, 1):
        sheet.column_dimensions[get_column_letter(index)].width = min(
            44, max(18, len(label) + 2)
        )
    _append(sheet, COLUMNS, header=True)
    totals = [Decimal("0")] * 4
    count = ordinary = retained = unavailable = undisbursed = 0
    for row in loan_directory_export_rows(
        workspace_id=workspace.pk, loan_filter=loan_filter
    ):
        count += 1
        if row["kind"] == "historical":
            retained += 1
            evidence, historical = row["evidence"], row["summary"]
            facts = evidence.document["facts"]
            opened_on = historical["opened_on"]
            values = (
                "Source record awaiting admission",
                facts.get("loan_number") or evidence.source_id,
                facts.get("borrower_name") or "Not recorded",
                "",
                "",
                "",
                "Closed claim",
                opened_on,
                completed_months(opened_on, as_of_date),
                None,
                None,
                None,
                None,
                None,
                "Unavailable: retained source claim is not financial debt",
                "Unavailable",
                "Unavailable: retained source claim; no verified operational collateral",
                *((None,) * (2 * len(CollateralMetal))),
            )
        else:
            ordinary += 1
            loan = row["loan"]
            position = get_loan_summary(loan, as_of_date=as_of_date)
            balance = position.balance
            if loan.state in {"DRAFT", "APPROVED"}:
                undisbursed += 1
            elif balance is None:
                unavailable += 1
            amounts = (
                (
                    balance.principal_outstanding,
                    balance.interest_outstanding,
                    balance.fees_outstanding,
                    balance.total_due,
                )
                if balance
                else (None,) * 4
            )
            if balance:
                totals = [total + value for total, value in zip(totals, amounts)]
            values = (
                "Preparation" if loan.loan_number is None else "Loan",
                loan.display_reference,
                loan.borrower.display_name,
                loan.borrower.party_code,
                loan.license.license_number,
                loan.series.pawn_display_name,
                loan.get_state_display(),
                loan.loan_date,
                position.age_months,
                position.unpaid_interest_months,
                *amounts,
                position.balance_basis,
                position.months_basis,
                *get_loan_collateral_summary(loan),
            )
        _append(
            sheet,
            tuple("Unavailable" if value is None else value for value in values),
            weight_columns_from=17,
        )
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{count + 1}"
    summary_sheet.column_dimensions["A"].width = 40
    summary_sheet.column_dimensions["B"].width = 100
    _append(summary_sheet, ("Filtered loan summary", "Value"), header=True)
    for values in (
        ("Workspace", workspace.name),
        ("As of", as_of_date),
        ("Matching records (all pages)", count),
        ("Loans", ordinary),
        ("Not yet disbursed", undisbursed),
        ("Source records awaiting admission", retained),
        ("Loans with unavailable balances", unavailable),
    ):
        _append(summary_sheet, values)
    for label, total in zip(
        (
            "Principal outstanding (INR)",
            "Interest pending (INR)",
            "Fees pending (INR)",
            "Total outstanding (INR)",
        ),
        totals,
    ):
        _append(summary_sheet, (label, "Unavailable" if unavailable else total))
        if unavailable:
            _append(summary_sheet, ("Known subtotal: " + label, total))
    for name, field in loan_filter.form.fields.items():
        value = loan_filter.form.cleaned_data.get(name)
        if value not in (None, ""):
            _append(summary_sheet, ("Filter: " + str(field.label), str(value)))
    _append(
        summary_sheet,
        (
            "Interest",
            "Current collection position includes eligible unposted interest; export does not post it. Legacy contracts retain recorded-debt semantics.",
        ),
    )
    _append(
        summary_sheet,
        (
            "Unpaid interest months",
            MONTHS_BASIS
            + " Reporting coverage, not saved payment allocation. Unavailable where monthly history is not established.",
        ),
    )
    _append(
        summary_sheet,
        (
            "Loan age",
            "Completed calendar months since the original loan date, with month-end clamping; independent of unpaid interest months.",
        ),
    )
    _append(
        summary_sheet,
        (
            "Collateral",
            "Summary lists all recorded items, piece quantities and current custody. Held weights are current grams in vault or with funding lender, excluding returned, disposed or transferred items; not original or historical stock. Item weights already cover the recorded quantity and are not multiplied by it. Missing measurements, absent collateral evidence and undisbursed records remain Unavailable. Unknown paper-closing handover makes that metal's held weights Unavailable. A known absent metal is zero. Missing gross never replaces known net.",
        ),
    )
    _append(
        summary_sheet,
        (
            "Scope",
            "Amounts reflect entered financial evidence. Paper-book completeness is not certified. Undisbursed and unadmitted source records do not enter debt totals. Khata accounts have their own register/export.",
        ),
    )
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
