"""Filtered Excel scope, financial evidence and reporting-month semantics."""
from datetime import date
from decimal import Decimal
from io import BytesIO
from unittest.mock import patch

from django.contrib.auth import get_user_model
from django.contrib.auth.models import Permission
from django.contrib.contenttypes.models import ContentType
from django.test import SimpleTestCase, override_settings
from django.urls import reverse
from openpyxl import load_workbook

from apps.orgs.models import Company, Membership, Role
from apps.tenancy.context import without_workspace_context, workspace_context
from apps.tenancy.testing import WorkspaceTestCase, workspace_role_permissions
from apps.tenant_apps.loans import models as m
from apps.tenant_apps.loans.selectors.loan_summary import completed_months, get_loan_summary, unpaid_month_count
from apps.tenant_apps.loans.selectors.party_loans import get_party_loan_history_summary
from apps.tenant_apps.loans.services.pawn_repayment import record_pawn_loan_repayment
from apps.tenant_apps.loans.services.pawn_reversal import reverse_pawn_loan_event
from apps.tenant_apps.party.models import Party
from . import test_unified_directory as directory_fixtures
from .test_continuation_admissions import AdmissionFixture


class MonthSummaryTests(SimpleTestCase):
    def test_age_is_completed_months_independent_of_unpaid_interest(self):
        for start, on, expected in ((date(2026, 1, 31), date(2026, 2, 28), 1),
                (date(2026, 1, 31), date(2026, 3, 30), 1),
                (date(2026, 1, 31), date(2026, 3, 31), 2),
                (date(2026, 5, 5), date(2026, 5, 5), 0),
                (None, date(2026, 5, 5), None)):
            self.assertEqual(completed_months(start, on), expected)

    def test_month_coverage_uses_each_actual_charge_and_counts_partial_month(self):
        # Oldest charges covered first: not debt / today's smaller charge.
        charges = [Decimal("100"), Decimal("80"), Decimal("60")]
        for pending, expected in ((0, 0), (1, 1), (60, 1), (61, 2), (140, 2), (141, 3), (241, None)):
            self.assertEqual(unpaid_month_count(charges, Decimal(pending)), expected)
        self.assertEqual(unpaid_month_count([0, 100, 0, 50], Decimal(51)), 2)


@override_settings(ROOT_URLCONF="django_project.workspace_urls")
class FilteredDirectoryExportTests(WorkspaceTestCase):
    setup_tenant = classmethod(directory_fixtures.UnifiedDirectoryTests.setup_tenant.__func__)
    _configured_setup = directory_fixtures.UnifiedDirectoryTests._configured_setup
    _payload = directory_fixtures.UnifiedDirectoryTests._payload
    accept = directory_fixtures.UnifiedDirectoryTests.accept

    def setUp(self):
        super().setUp()
        self.owner = self.tenant.owner
        self.start_active_trial()
        self.client = self.make_workspace_client()
        self.client.force_login(self.owner)
        self.party = Party.objects.create(display_name="Fictional borrower")
        from apps.tenant_apps.loans.services.product_catalog import _seed_default_loan_products
        self.product_version = _seed_default_loan_products()[0]
        type(self.product_version).objects.filter(pk=self.product_version.pk).update(status="ACTIVE")
        self.licence, self.series = self._configured_setup()
        self.url = reverse("loans:pawn_loan_list")

    def make_loan(self, number, **extra):
        return m.PawnLoan.objects.create(workspace=self.tenant, borrower=self.party,
            license=self.licence, series=self.series, product_version=self.product_version,
            loan_number=number, state="DRAFT", principal_amount=100, monthly_interest_rate=2,
            tenure_months=12, loan_date=date(2026, 5, 5), **extra)

    def export(self, **query):
        response = self.client.get(self.url, {**query, "export": "xlsx"})
        self.assertEqual(response.status_code, 200)
        self.assertIn("no-store", response["Cache-Control"])
        return load_workbook(BytesIO(response.content), data_only=False)

    def test_export_matches_all_filtered_pages_and_retains_unknown_source_claims(self):
        for number in range(27):
            self.make_loan(f"MATCH-{number:02}")
        self.make_loan("UNRELATED")
        self.accept(key="MATCH-SOURCE", name="=Fictional customer", amount=None)
        before = (m.PawnLoan.objects.count(), m.PawnLoanEvent.objects.count(), m.HistoricalLoanEvidence.objects.count())
        page = self.client.get(self.url, {"q": "MATCH", "page": 2})
        self.assertEqual(page.context["page_obj"].paginator.count, 28)
        self.assertContains(page, "export=xlsx")
        book = self.export(q="MATCH", page=2)
        rows = list(book["Loans"].values)
        self.assertEqual(len(rows), 29)
        self.assertEqual(len([row for row in rows[1:] if row[0] == "Loan"]), 27)
        source = rows[-1]
        self.assertEqual(source[2], "=Fictional customer")
        self.assertEqual(book["Loans"]["C29"].data_type, "s")
        self.assertEqual(source[10:14], ("Unavailable",) * 4)
        summary = dict(list(book["Summary"].values)[1:])
        self.assertEqual(summary["Matching records (all pages)"], 28)
        self.assertEqual(summary["Total outstanding (INR)"], 0)
        self.assertEqual(before, (m.PawnLoan.objects.count(), m.PawnLoanEvent.objects.count(), m.HistoricalLoanEvidence.objects.count()))

    def test_borrower_state_licence_series_and_dates_use_the_screen_filters(self):
        first = self.make_loan("SELECTED")
        self.make_loan("OUTSIDE-DATE")
        m.PawnLoan.objects.filter(loan_number="OUTSIDE-DATE").update(loan_date=date(2026, 4, 1))
        query = dict(borrower=self.party.pk, state="DRAFT", license=self.licence.pk,
            series=self.series.pk, loan_date_from="2026-05-01", loan_date_to="2026-05-31", records="ordinary")
        self.assertEqual(self.client.get(self.url, query).context["loans"], [first])
        self.assertEqual(list(self.export(**query)["Loans"].values)[1][1], "SELECTED")
        self.assertEqual(len(list(self.export(records="historical")["Loans"].values)), 1)
        for invalid in ({"state": "INVALID"}, {"loan_date_from": "bad"},
                {"loan_date_from": "2026-06-01", "loan_date_to": "2026-05-01"}, {"borrower": 999999999}):
            self.assertEqual(self.client.get(self.url, {**invalid, "export": "xlsx"}).status_code, 400)

    def test_collateral_export_separates_metals_and_excludes_returned_weights(self):
        loan = self.make_loan("COLLATERAL")
        m.PawnLoan.objects.filter(pk=loan.pk).update(state="ACTIVE")

        def item(description, metal, gross, net, quantity=1, custody="IN_VAULT"):
            return m.PawnCollateralItem.objects.create(workspace=self.tenant, loan=loan,
                description=description, metal=metal, gross_weight=gross, net_weight=net,
                purity_percentage=90, quantity=quantity, custody_state=custody)

        item("=Gold bangles", "GOLD", "12.3456", "10.1234", quantity=2)
        item("Chain with lender", "GOLD", "6.1111", "5.1111", custody="WITH_FUNDING_LENDER")
        item("Returned ring", "GOLD", "100", "90", custody="WITH_CUSTOMER")
        item("Silver anklets", "SILVER", "80", "75", quantity=2)
        item("Bronze vessel", "BRONZE", None, "200.9876", quantity=None)
        item("Other article", "OTHER", "15", "14")
        self.make_loan("OUTSIDE-COLLATERAL-FILTER")
        book = self.export(borrower=self.party.pk, state="ACTIVE")
        sheet = book["Loans"]
        rows = list(sheet.values)
        self.assertEqual(len(rows), 2)
        data = dict(zip(rows[0], rows[1]))
        self.assertEqual(data["Loan"], "COLLATERAL")
        # Weights belong to an entire item row, even when quantity > 1.
        self.assertEqual(Decimal(str(data["Held gold gross weight (g)"])), Decimal("18.4567"))
        self.assertEqual(Decimal(str(data["Held gold net weight (g)"])), Decimal("15.2345"))
        self.assertEqual(data["Held silver net weight (g)"], 75)
        self.assertEqual(data["Held bronze gross weight (g)"], "Unavailable")
        self.assertEqual(Decimal(str(data["Held bronze net weight (g)"])), Decimal("200.9876"))
        self.assertEqual(data["Held other net weight (g)"], 14)
        self.assertIn("=Gold bangles; Gold; qty 2; In Vault", data["Collateral summary"])
        self.assertIn("Returned ring", data["Collateral summary"])
        self.assertIn("With Customer", data["Collateral summary"])
        self.assertIn("qty not recorded", data["Collateral summary"])
        self.assertEqual(sheet["Q2"].data_type, "s")
        self.assertEqual(sheet["R2"].number_format, "#,##0.0000")
        self.assertEqual(sheet["K2"].number_format, "General")  # Unavailable balance.
        self.assertEqual(sheet.auto_filter.ref, "A1:Y2")

    def test_missing_gross_and_released_holdings_remain_distinct(self):
        loan = self.make_loan("WEIGHT-UNKNOWN")
        m.PawnLoan.objects.filter(pk=loan.pk).update(state="ACTIVE")
        for gross, net in ((None, "2"), ("4", "3")):
            m.PawnCollateralItem.objects.create(workspace=self.tenant, loan=loan,
                description="Chain", metal="GOLD", gross_weight=gross, net_weight=net,
                purity_percentage=90)
        rows = list(self.export(q="WEIGHT-UNKNOWN")["Loans"].values)
        data = dict(zip(rows[0], rows[1]))
        self.assertEqual(data["Held gold gross weight (g)"], "Unavailable")
        self.assertEqual(data["Held gold net weight (g)"], 5)
        self.assertEqual(data["Held silver net weight (g)"], 0)
        # Exercise the read projection after a completed custody transition.
        m.PawnCollateralItem.objects.filter(loan=loan).update(custody_state="WITH_CUSTOMER")
        rows = list(self.export(q="WEIGHT-UNKNOWN")["Loans"].values)
        data = dict(zip(rows[0], rows[1]))
        self.assertEqual(data["Held gold gross weight (g)"], 0)
        self.assertEqual(data["Held gold net weight (g)"], 0)
        self.assertIn("With Customer", data["Collateral summary"])
        m.PawnCollateralItem.objects.filter(loan=loan).update(custody_state="PAPER_CLOSED")
        rows = list(self.export(q="WEIGHT-UNKNOWN")["Loans"].values)
        data = dict(zip(rows[0], rows[1]))
        self.assertEqual(data["Held gold gross weight (g)"], "Unavailable")
        self.assertEqual(data["Held gold net weight (g)"], "Unavailable")
        self.assertEqual(data["Held silver net weight (g)"], 0)

    def test_no_collateral_evidence_and_undisbursed_records_do_not_claim_zero_holdings(self):
        empty = self.make_loan("EMPTY-COLLATERAL")
        m.PawnLoan.objects.filter(pk=empty.pk).update(state="ACTIVE")
        draft = self.make_loan("DRAFT-COLLATERAL")
        m.PawnCollateralItem.objects.create(workspace=self.tenant, loan=draft,
            description="Draft ring", metal="GOLD", gross_weight=3, net_weight=2,
            purity_percentage=90)
        self.accept(key="SOURCE-COLLATERAL", amount=None)
        rows = list(self.export()["Loans"].values)
        for row in rows[1:]:
            self.assertEqual(row[17:], ("Unavailable",) * 8)
        records = {row[1]: row for row in rows[1:]}
        self.assertEqual(records[empty.loan_number][16], "Not recorded")
        self.assertIn("Draft ring", records[draft.loan_number][16])

    def test_collateral_previews_are_compact_escaped_and_borrower_scoped_on_both_lists(self):
        active = self.make_loan("PREVIEW-ACTIVE")
        closed = self.make_loan("PREVIEW-CLOSED")
        m.PawnLoan.objects.filter(pk=active.pk).update(state="ACTIVE")
        m.PawnLoan.objects.filter(pk=closed.pk).update(state="CLOSED")
        for loan, description, metal, custody in (
                (active, "<script>Ring</script>", "GOLD", "IN_VAULT"),
                (active, "Silver anklets", "SILVER", "WITH_FUNDING_LENDER"),
                (active, "Third item beyond compact preview", "BRONZE", "IN_VAULT"),
                (closed, "Returned chain", "GOLD", "WITH_CUSTOMER")):
            m.PawnCollateralItem.objects.create(workspace=self.tenant, loan=loan,
                description=description, metal=metal, gross_weight=3, net_weight="2.1234",
                purity_percentage=90, quantity=None if loan == closed else 2, custody_state=custody)
        other = Party.objects.create(display_name="Other preview borrower")
        foreign = self.make_loan("OTHER-PREVIEW-LOAN")
        m.PawnLoan.objects.filter(pk=foreign.pk).update(borrower=other)
        m.PawnCollateralItem.objects.create(workspace=self.tenant, loan=foreign,
            description="OTHER-BORROWER-COLLATERAL", metal="GOLD", gross_weight=3,
            net_weight=2, purity_percentage=90)
        before = m.PawnLoanEvent.objects.count()
        responses = [self.client.get(self.url, {"borrower": self.party.pk}),
            self.client.get(self.url, {"borrower": self.party.pk},
                HTTP_HX_REQUEST="true", HTTP_HX_TARGET="loan-results"),
            self.client.get(reverse("party:party_detail", args=[self.party.pk]), {"tab": "loans"})]
        for response in responses:
            self.assertEqual(response.status_code, 200)
            self.assertContains(response, "data-collateral-summary", count=2)
            self.assertContains(response, "&lt;script&gt;Ring&lt;/script&gt;")
            self.assertNotContains(response, "<script>Ring</script>")
            self.assertContains(response, "Gold 2.1234 g net")
            self.assertContains(response, "Silver 2.1234 g net")
            self.assertContains(response, "qty 2")
            self.assertContains(response, "+ 1 more item entry")
            self.assertNotContains(response, "Third item beyond compact preview")
            self.assertContains(response, "Returned chain")
            self.assertContains(response, "With Customer")
            self.assertContains(response, "qty not recorded")
            self.assertNotContains(response, "OTHER-BORROWER-COLLATERAL")
        m.PawnCollateralItem.objects.filter(loan=closed).update(custody_state="PAPER_CLOSED")
        response = self.client.get(reverse("party:party_detail", args=[self.party.pk]), {"tab": "loans"})
        self.assertContains(response, "Handover unconfirmed")
        m.PawnLoan.objects.filter(pk=active.pk).update(state="DRAFT")
        response = self.client.get(self.url, {"q": active.loan_number})
        self.assertContains(response, "Silver anklets")
        self.assertNotContains(response, "In Vault")
        self.assertNotContains(response, "With Funding Lender")
        self.assertEqual(m.PawnLoanEvent.objects.count(), before)

    def test_foreign_workspace_records_and_filter_ids_are_not_exported(self):
        from uuid import uuid4
        from apps.tenant_apps.data_portability.tests.test_loan_archive import document
        from apps.tenant_apps.loans.services.archive import accept_evidence
        from apps.tenant_apps.loans.services.history_contract import digest
        other = Company.objects.create(name="Export isolation", schema_name=uuid4().hex,
            owner=self.owner, creator=self.owner)
        value = document()
        value["source"]["loan_id"] = "FOREIGN-EXPORT-SECRET"
        value["facts"]["loan_number"] = "FOREIGN-EXPORT-SECRET"
        with without_workspace_context(), workspace_context(other.pk):
            Membership.objects.get_or_create(company=other, user=self.owner, defaults={"role": Role.objects.get(name="Owner")})
            foreign_party = Party.objects.create(display_name="Foreign customer")
            accept_evidence(workspace_id=other.pk, actor=self.owner, document=value,
                expected_sha256=digest(value), confirmed=True)
        book = self.export(q="FOREIGN-EXPORT-SECRET")
        self.assertEqual(len(list(book["Loans"].values)), 1)
        self.assertEqual(self.client.get(self.url, {"export": "xlsx", "borrower": foreign_party.pk}).status_code, 400)

    def test_export_requires_report_permission_in_addition_to_view_and_is_read_only(self):
        role = Role.objects.create(name="Filtered report reader")
        user = get_user_model().objects.create_user(username="filtered-reader")
        ct = ContentType.objects.get_for_model(Company)
        view, _ = Permission.objects.get_or_create(content_type=ct, codename="data_view", defaults={"name": "View"})
        export, _ = Permission.objects.get_or_create(content_type=ct, codename="report_export", defaults={"name": "Export"})
        workspace_role_permissions(role, self.tenant).add(view)
        Membership.objects.create(company=self.tenant, user=user, role=role)
        self.client.force_login(user)
        self.assertNotContains(self.client.get(self.url), "data-loan-export")
        self.assertEqual(self.client.get(self.url, {"export": "xlsx"}).status_code, 403)
        workspace_role_permissions(role, self.tenant).add(export)
        self.assertContains(self.client.get(self.url), "data-loan-export")
        self.export()
        self.assertEqual(self.client.post(self.url, {"export": "xlsx"}).status_code, 405)
        self.assertEqual(self.client.get(self.url, {"export": "csv"}).status_code, 400)


class FinancialSummaryTests(AdmissionFixture):
    def test_four_origins_share_collection_amounts_without_posting_and_keep_month_unknowns(self):
        day = date(2026, 7, 6)
        with self.scoped(), patch("django.utils.timezone.localdate", return_value=day):
            before = list(m.PawnLoanEvent.objects.values_list("pk", "payload_fingerprint"))
            for loan in self.loans:
                summary = get_loan_summary(loan, as_of_date=day)
                self.assertEqual(summary.age_months, 3)
                self.assertEqual(summary.balance.principal_outstanding, self.remaining)
                self.assertEqual(summary.balance.interest_outstanding, self.monthly * 3)
                self.assertEqual(summary.unpaid_interest_months, None if loan == self.opening else 3)
                self.assertEqual(summary.balance.total_due, self.remaining + self.monthly * 3)
            history = get_party_loan_history_summary(self.direct.borrower, limit=1)
            owned = sum(loan.borrower_id == self.direct.borrower_id for loan in self.loans)
            self.assertEqual(history["counts"]["active_principal"], self.remaining * owned)
            self.assertEqual(history["counts"]["active_interest"], self.monthly * 3 * owned)
            self.assertEqual(history["counts"]["active_outstanding"], (self.remaining + self.monthly * 3) * owned)
            self.assertEqual(before, list(m.PawnLoanEvent.objects.values_list("pk", "payload_fingerprint")))
            from apps.tenant_apps.loans.filters import PawnLoanFilter
            from apps.tenant_apps.loans.services.directory_export import export_loan_directory
            filters = PawnLoanFilter({"state": "ACTIVE"},
                queryset=m.PawnLoan.objects.filter(workspace=self.a), workspace=self.a)
            book = load_workbook(BytesIO(export_loan_directory(workspace=self.a, loan_filter=filters, as_of_date=day)))
            rows = list(book["Loans"].values)[1:]
            self.assertEqual(len(rows), 4)
            for row in rows:
                self.assertEqual(Decimal(str(row[10])), self.remaining)
                self.assertEqual(Decimal(str(row[11])), self.monthly * 3)
                self.assertEqual(Decimal(str(row[12])), 0)
                self.assertEqual(Decimal(str(row[13])), self.remaining + self.monthly * 3)
                self.assertEqual(row[8], 3)
            summary = dict(list(book["Summary"].values)[1:])
            self.assertEqual(Decimal(str(summary["Total outstanding (INR)"])), (self.remaining + self.monthly * 3) * 4)
            self.assertEqual(before, list(m.PawnLoanEvent.objects.values_list("pk", "payload_fingerprint")))

    def test_partial_interest_receipt_and_reversal_change_unpaid_months_using_saved_charges(self):
        day = date(2026, 7, 6)
        with self.scoped(), patch("django.utils.timezone.localdate", return_value=day):
            for loan in (self.direct, self.paper, self.imported):
                receipt = record_pawn_loan_repayment(loan.pk, amount=self.monthly + 1,
                    request_key=f"summary-interest-{loan.pk}", actor=self.actor)
                self.assertEqual(get_loan_summary(loan, as_of_date=day).unpaid_interest_months, 2)
                if loan == self.direct:
                    reverse_pawn_loan_event(receipt.loan_event.pk, reason="Synthetic receipt correction", actor=self.actor)
                    self.assertEqual(get_loan_summary(loan, as_of_date=day).unpaid_interest_months, 3)

    def test_unavailable_active_balance_prevents_misleading_partial_totals(self):
        with self.scoped():
            m.PawnLoan.objects.create(workspace=self.a, borrower=self.direct.borrower,
                license=self.direct.license, series=self.direct.series, product_version=self.direct.product_version,
                loan_number="MISSING-SUMMARY-ORIGIN", loan_date=self.original, principal_amount=100,
                monthly_interest_rate=2, tenure_months=12, state="ACTIVE")
            totals = get_party_loan_history_summary(self.direct.borrower)["counts"]
            for field in ("active_outstanding", "active_principal", "active_interest", "active_fees"):
                self.assertIsNone(totals[field])
